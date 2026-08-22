"""SBA Active Mentor-Protege Agreements (All Small MPP) - enrichment, no key.

WHY: the relationship edge itself. Each row is an ACTIVE mentor-protege
agreement carrying approval date, agreement NAICS, protege cert types
(for example "EDWOSB,WOSB,8(a)"), and the UEI for BOTH sides. No pipeline
source (SAM, USAspending, notices) carries mentor-protege agreement
existence, scope NAICS, or approval date, so this is first-class teaming
evidence: who has formally committed to whom, in which NAICS, since when.

Document page: https://www.sba.gov/document/support--active-mentor-protege-agreements
  The page 302s from www.sba.gov to legacy.sba.gov mid-site-migration, and
  the xlsx filename churns per refresh (it carries an export counter and an
  effective date), so every pull scrapes the page for the current .xlsx
  link. The file itself resolves ONLY on legacy.sba.gov: the same path on
  www.sba.gov returns 404 (verified 2026-08-08), so relative hrefs resolve
  against the legacy host, never the page's own canonical URL.

File quirks (manual pivot export, guarded here):
  - two preamble rows and one blank row before the header: the header row is
    located by scanning, never by a fixed offset;
  - a trailing "Sum of Count" pivot column: dropped (never mapped);
  - malformed UEIs as published (13-char, 11-char, empty, even a company
    name in the UEI column): nulled, while the pair row is kept.

Snapshot-only publication: terminated or expired agreements vanish silently
between refreshes, so a pull is a census of ACTIVE agreements as of the
page's stated effective date, never a history.
"""

from __future__ import annotations

import io
import re
import time
import unicodedata
from datetime import date, datetime, timezone
from html import unescape
from typing import Any, Callable, Optional

from tools.api import _http
from tools.api.base import (DataSource, SourceKind, SourceQuery,
                            cap_disclosed, register_source)
from tools.api.provenance import make_provenance_envelope

DOC_PAGE_URL = "https://www.sba.gov/document/support--active-mentor-protege-agreements"
#: relative hrefs resolve here; www.sba.gov 404s the file path (2026-08-08)
LEGACY_BASE = "https://legacy.sba.gov"
#: hard ceiling on returned pairs; cap_disclosed says any cut out loud
PAIR_CAP = 200
#: header-scan window; the pivot export has 3 junk rows today, headroom for churn
HEADER_SCAN_ROWS = 12

_XLSX_HREF = re.compile(r'href="([^"]+\.xlsx[^"]*)"', re.IGNORECASE)
_EFFECTIVE = re.compile(r"Effective</strong>:\s*([A-Za-z]+ \d{1,2}, \d{4})")
_UEI = re.compile(r"^[A-Z0-9]{12}$")

#: normalized header text -> output field; "sum of count" is absent on purpose
_COLUMNS = {
    "date approved": "approved",
    "naics": "naics",
    "protege": "protege",
    "protege unique entity id": "protege_uei",
    "mentor": "mentor",
    "mentor unique entity id": "mentor_uei",
    "cert type": "cert_types",
}


def _get_bytes(url: str, *, timeout: float = 60.0, retries: int = 3) -> bytes:
    """Binary GET on the shared _http client (IPv4-forced transport, redirects).

    _http exposes json/text helpers only and an xlsx needs raw bytes, so this
    reuses its client factory and mirrors its bounded-retry ladder. Private to
    this adapter; promote a get_bytes helper into _http if a second binary
    source lands.
    """
    last_exc: Optional[Exception] = None
    for attempt in range(retries):
        try:
            with _http._client(timeout) as client:
                resp = client.get(url)
        except Exception as exc:  # noqa: BLE001 - transport errors retry
            last_exc = exc
        else:
            if resp.status_code in _http._RETRYABLE:
                last_exc = RuntimeError(
                    f"HTTP {resp.status_code} fetching {url}")
            elif resp.status_code >= 400:
                # non-retryable status fails identically on every attempt
                raise RuntimeError(f"HTTP {resp.status_code} fetching {url}")
            else:
                return resp.content
        if attempt < retries - 1:
            time.sleep(0.5 * (2 ** attempt))
    assert last_exc is not None
    raise last_exc


def _xlsx_url_from_page(html: str) -> Optional[str]:
    """First .xlsx href on the document page, made absolute on the legacy host."""
    match = _XLSX_HREF.search(html or "")
    if not match:
        return None
    href = unescape(match.group(1)).strip()
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if not href.startswith("/"):
        href = "/" + href
    return LEGACY_BASE + href


def _effective_date(html: str) -> Optional[str]:
    """The page's own 'Effective' date (the snapshot's data_as_of), ISO format."""
    match = _EFFECTIVE.search(html or "")
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%B %d, %Y").date().isoformat()
    except ValueError:
        return None


def _grid_from_xlsx(data: bytes) -> list:
    """First worksheet as a row grid (streaming read, values only)."""
    import openpyxl  # local: only the live path pays the import

    workbook = openpyxl.load_workbook(
        io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _norm_header(value: Any) -> str:
    """Accent- and punctuation-proof header key ('Protégé' -> 'protege')."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _clean_str(value: Any) -> Optional[str]:
    text = str(value).strip() if value is not None else ""
    return text or None


def _valid_uei(value: Any) -> Optional[str]:
    """12-char alphanumeric or None; the file publishes several malformed ones."""
    text = str(value).strip().upper() if value is not None else ""
    return text if _UEI.match(text) else None


def _iso_date(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip() if value is not None else ""
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return text  # keep the source's own string rather than losing it


def _naics_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if text.endswith(".0"):  # float leakage in pivot exports
        text = text[:-2]
    return text or None


def _pairs_from_grid(grid: list) -> "tuple[list, int]":
    """Normalize the pivot-export grid -> (pairs, dropped_row_count).

    Scans for the header row (the preamble is 2 text rows plus 1 blank today,
    but the count is not trusted), maps columns by normalized name so the
    trailing 'Sum of Count' pivot column falls away, nulls invalid UEIs, and
    drops-and-counts rows with content but no mentor or protege name. Blank
    filler rows are skipped silently: structure, not records.
    """
    header_at: Optional[int] = None
    colmap: dict = {}
    for index, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        normed = [_norm_header(cell) for cell in (row or [])]
        if {"date approved", "protege", "mentor"}.issubset(set(normed)):
            header_at = index
            for position, name in enumerate(normed):
                field = _COLUMNS.get(name)
                if field is not None and field not in colmap:
                    colmap[field] = position
            break
    if header_at is None:
        raise ValueError(
            f"header row not found in first {HEADER_SCAN_ROWS} rows; "
            "the pivot-export layout changed")

    pairs: list = []
    dropped = 0
    for row in grid[header_at + 1:]:
        cells = list(row) if row is not None else []
        if all(cell is None or str(cell).strip() == "" for cell in cells):
            continue

        def _cell(field: str, cells: list = cells) -> Any:
            position = colmap.get(field)
            if position is None or position >= len(cells):
                return None
            return cells[position]

        mentor = _clean_str(_cell("mentor"))
        protege = _clean_str(_cell("protege"))
        if mentor is None and protege is None:
            dropped += 1  # content without either party: not a usable edge
            continue
        pairs.append({
            "mentor": mentor,
            "mentor_uei": _valid_uei(_cell("mentor_uei")),
            "protege": protege,
            "protege_uei": _valid_uei(_cell("protege_uei")),
            "approved": _iso_date(_cell("approved")),
            "naics": _naics_str(_cell("naics")),
            "cert_types": _clean_str(_cell("cert_types")),
        })
    return pairs, dropped


def _matches(pair: dict, keywords: list, naics_codes: list,
             approved_from: Optional[str], approved_to: Optional[str]) -> bool:
    """Keywords hit mentor OR protege name; NAICS is prefix-scoped; the
    query's posted window translates onto the approval date."""
    if keywords:
        hay = " ".join(
            value for value in (pair.get("mentor"), pair.get("protege")) if value
        ).lower()
        if not any(keyword in hay for keyword in keywords):
            return False
    if naics_codes:
        naics = pair.get("naics") or ""
        if not any(naics.startswith(code) for code in naics_codes):
            return False
    if approved_from or approved_to:
        approved = pair.get("approved")
        if not approved:
            return False  # cannot place it in the window: excluded, not guessed
        if approved_from and approved < approved_from:
            return False
        if approved_to and approved > approved_to:
            return False
    return True


def _register_when_cataloged(cls: type) -> type:
    """@register_source, tolerated until the SourceSpec lands.

    register_source refuses uncataloged adapters by design. The catalog entry
    and the registrar import are wired by the integration lead; until then
    this module must stay importable and testable on its own, so the refusal
    downgrades to a no-op. The moment the SourceSpec exists, importing this
    module registers the adapter exactly like every other source.
    """
    try:
        return register_source(cls)
    except ValueError:
        return cls


@_register_when_cataloged
class SbaMentorProtegeSource(DataSource):
    name = "sba_mentor_protege"
    kind = SourceKind.ENRICHMENT

    def __init__(self,
                 fetch_text: Optional[Callable[..., str]] = None,
                 fetch_bytes: Optional[Callable[..., bytes]] = None) -> None:
        self._fetch_text = fetch_text
        self._fetch_bytes = fetch_bytes

    # ------------------------------------------------------------------ #
    def search(self, query: SourceQuery) -> list:
        raise NotImplementedError(
            "enrichment-only source; use enrich() or pairs_for()")

    def healthcheck(self) -> "tuple[bool, str]":
        """Probe the document page AND that it still lists an .xlsx link."""
        try:
            fetch = self._fetch_text or _http.get_text
            html = fetch(DOC_PAGE_URL, timeout=8.0, retries=1)
            url = _xlsx_url_from_page(html)
            if not url:
                return False, "document page reachable but no .xlsx link found"
            from urllib.parse import unquote
            return True, f"document page lists {unquote(url.rsplit('/', 1)[-1])}"
        except Exception as exc:  # noqa: BLE001
            return False, f"document page unreachable: {exc}"

    # ------------------------------------------------------------------ #
    def enrich(self, query: SourceQuery) -> dict:
        """Active mentor-protege pairs scoped to the query.

        Returns {"pairs": [{mentor, mentor_uei, protege, protege_uei,
        approved, naics, cert_types}], "total_matched", "total_agreements",
        "dropped_rows", "document_page", "xlsx_url", "_provenance"} with any
        cap said out loud via cap_disclosed. Never raises: failures come
        back as {"pairs": [], "error", "_provenance": {...status failed...}}.
        """
        retrieved_at = datetime.now(timezone.utc)
        try:
            fetch_text = self._fetch_text or _http.get_text
            html = fetch_text(DOC_PAGE_URL)
            xlsx_url = _xlsx_url_from_page(html)
            if not xlsx_url:
                raise RuntimeError(
                    "no .xlsx link on the document page; layout changed")
            fetch_bytes = self._fetch_bytes or _get_bytes
            grid = _grid_from_xlsx(fetch_bytes(xlsx_url))
            pairs, dropped = _pairs_from_grid(grid)
        except Exception as exc:  # noqa: BLE001
            return self._failure_payload(exc, retrieved_at)

        keywords = [k.lower().strip('"').strip() for k in (query.keywords or [])
                    if k and k.strip()]
        naics_codes = [str(code).strip() for code in (query.naics_codes or [])
                       if str(code).strip()]
        approved_from = query.posted_from.isoformat() if query.posted_from else None
        approved_to = query.posted_to.isoformat() if query.posted_to else None
        matched = [pair for pair in pairs
                   if _matches(pair, keywords, naics_codes,
                               approved_from, approved_to)]

        cap = max(1, min(PAIR_CAP, query.limit or PAIR_CAP))
        out = cap_disclosed(matched, cap, key="pairs")
        out["total_agreements"] = len(pairs)
        out["dropped_rows"] = dropped
        out["document_page"] = DOC_PAGE_URL
        out["xlsx_url"] = xlsx_url

        effective = _effective_date(html)
        limitations = [
            "snapshot-only publication: terminated or expired agreements "
            "vanish silently between refreshes",
        ]
        if dropped:
            limitations.append(
                f"{dropped} rows dropped: neither mentor nor protege name")
        if effective is None:
            limitations.append("effective date not stated on the document page")
        out["_provenance"] = make_provenance_envelope(
            self.name,
            status="complete",
            mode="live_page_scrape_xlsx",
            retrieval_mode="live",
            retrieved_at=retrieved_at,
            data_as_of=effective,
            record_count=len(out["pairs"]),
            limitations=limitations,
        ).model_dump(mode="json")
        return out

    def pairs_for(self, query: SourceQuery) -> list:
        """The matched pairs alone, for engines that join rather than render."""
        return self.enrich(query).get("pairs", [])

    # ------------------------------------------------------------------ #
    def _failure_payload(self, exc: Exception, retrieved_at: datetime) -> dict:
        envelope = make_provenance_envelope(
            self.name,
            status="failed",
            mode="live_page_scrape_xlsx",
            retrieval_mode="live",
            retrieved_at=retrieved_at,
            record_count=0,
            limitations=["fetch or parse failed; no pairs available this pull"],
        )
        return {
            "pairs": [],
            "total_matched": 0,
            "error": str(exc),
            "document_page": DOC_PAGE_URL,
            "_provenance": envelope.model_dump(mode="json"),
        }
