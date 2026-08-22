"""OMB SF-133 monthly budget-execution workbooks (official MAX portal).

SF-133 lines describe budget resources, obligations, outlays, and balances at
Treasury-account level.  They are account context, never solicitations,
contract awards, or a claim that a displayed balance will be procured.
"""

from __future__ import annotations

import hashlib
import io
import re
from datetime import date, datetime
from typing import Any, Optional

from openpyxl import load_workbook

from tools.api._accepted_wave import (FetchBytes, fetch_bytes, payload,
                                      receipt)
from tools.api.base import DataSource, SourceKind, SourceQuery, register_source


SOURCE_NAME = "omb_sf133"
MODE = "live_omb_sf133_monthly_xlsx"
ATTACHMENT_ROOT = (
    "https://portal.max.gov/portal/document/SF133/Budget/attachments/2692285459"
)
FISCAL_YEAR = 2026
AGENCY_SLUGS = {
    "dhs": "Department_of_Homeland_Security",
    "department of homeland security": "Department_of_Homeland_Security",
    "homeland security": "Department_of_Homeland_Security",
    "doe": "Department_of_Energy",
    "department of energy": "Department_of_Energy",
    "energy": "Department_of_Energy",
    "hhs": "Department_of_Health_and_Human_Services",
    "department of health and human services": "Department_of_Health_and_Human_Services",
    "health and human services": "Department_of_Health_and_Human_Services",
    "sba": "Small_Business_Administration",
    "small business administration": "Small_Business_Administration",
}
DROP_FIELDS = {"F2_USER_ID", "F2_USER"}
AMOUNT_PREFIX = "AMT"


def _url(slug: str, fiscal_year: int = FISCAL_YEAR) -> str:
    return (
        f"{ATTACHMENT_ROOT}/FY{fiscal_year}_SF133_MONTHLY_{slug}.xlsx"
    )


def _clean(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return " ".join(value.split()) or None
    return value


def _code(value: Any, width: int) -> Optional[str]:
    text = str(value or "").strip()
    return text.zfill(width) if text else None


def _agency_slug(query: SourceQuery) -> tuple[Optional[str], list[str]]:
    limitations: list[str] = []
    agencies = [str(value).strip() for value in query.agencies if str(value).strip()]
    if not agencies:
        return None, ["unscoped retrieval refused; pass one supported agency"]
    def aliases(agency: str) -> tuple[str, ...]:
        full = " ".join(agency.casefold().split())
        without_parenthetical = re.sub(r"\s*\([^)]*\)\s*$", "", full)
        parenthetical = re.findall(r"\(([^)]*)\)", full)
        return tuple(dict.fromkeys(
            [full, without_parenthetical]
            + [" ".join(value.split()) for value in parenthetical]
        ))

    selected_index = next(
        (index for index, agency in enumerate(agencies)
         if any(alias in AGENCY_SLUGS for alias in aliases(agency))), None)
    if selected_index is None:
        return None, [
            f"no verified FY{FISCAL_YEAR} workbook URL mapping for requested agencies"
        ]
    selected = agencies[selected_index]
    skipped = agencies[:selected_index]
    if skipped:
        limitations.append(
            "skipped unsupported agency mappings before first supported term: "
            + ", ".join(skipped))
    if len(agencies) > 1:
        limitations.append(
            f"bounded collector used supported agency {selected!r} from "
            f"{len(agencies)} requested terms")
    selected_alias = next(alias for alias in aliases(selected)
                          if alias in AGENCY_SLUGS)
    return AGENCY_SLUGS[selected_alias], limitations


def _stable_id(row: dict[str, Any]) -> str:
    basis = "|".join(str(row.get(key) or "") for key in (
        "TAFS", "LINENO", "CAT_B", "STAT", "COHORT", "SECTION_NO",
    ))
    return "sf133:" + hashlib.sha256(basis.encode("utf-8")).hexdigest()[:24]


def _matches(row: dict[str, Any], query: SourceQuery) -> bool:
    terms = [str(value).strip().strip('"').casefold()
             for value in query.keywords if str(value).strip()]
    if not terms:
        return True
    haystack = " ".join(str(row.get(key) or "") for key in (
        "LINE_DESC", "LINE_DESC_SHORT", "CAT_B", "PGM_CAT", "OMB_ACCOUNT",
        "AGENCY_TITLE", "BUREAU_TITLE", "TAFS", "OMB_ACCT",
    )).casefold()
    return any(term in haystack for term in terms)


def _parse_workbook(blob: bytes, query: SourceQuery) -> tuple[list[dict[str, Any]], int]:
    if not blob.startswith(b"PK"):
        raise ValueError("SF-133 response is not an XLSX archive")
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if "Raw Data" not in workbook.sheetnames:
        raise ValueError("SF-133 workbook has no Raw Data sheet")
    sheet = workbook["Raw Data"]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(value or "").strip().upper() for value in next(rows)]
    except StopIteration as exc:
        raise ValueError("SF-133 Raw Data sheet is empty") from exc
    required = {"TAFS", "LINENO", "LINE_DESC", "AGENCY_TITLE", "LAST_UPDATED"}
    if not required.issubset(header):
        raise ValueError(
            "SF-133 Raw Data header missing " + ", ".join(sorted(required - set(header)))
        )

    normalized: list[dict[str, Any]] = []
    dropped = 0
    for values in rows:
        native = {key: _clean(value) for key, value in zip(header, values)
                  if key and key not in DROP_FIELDS}
        if not native.get("LINE_DESC") or not native.get("TAFS"):
            dropped += 1
            continue
        if not _matches(native, query):
            continue
        amounts = {
            key: value for key, value in native.items()
            if key.startswith(AMOUNT_PREFIX) and isinstance(value, (int, float))
            and not isinstance(value, bool)
        }
        row = {
            "record_id": _stable_id(native),
            "fiscal_year": native.get("RPT_YR"),
            "agency_code": native.get("AGENCY"),
            "agency": native.get("AGENCY_TITLE"),
            "bureau_code": native.get("BUREAU"),
            "bureau": native.get("BUREAU_TITLE"),
            "tafs": native.get("TAFS"),
            "treasury_agency": native.get("TRAG"),
            "main_account_code": native.get("TRACCT"),
            "availability_start_year": native.get("FY1"),
            "availability_end_year": native.get("FY2"),
            "account_status": native.get("STAT"),
            "omb_account": native.get("OMB_ACCT"),
            "omb_account_title": native.get("OMB_ACCOUNT"),
            "line_number": native.get("LINENO"),
            "line_description": native.get("LINE_DESC"),
            "category_b": native.get("CAT_B"),
            "section": native.get("SECTION"),
            "amounts_by_report_column": amounts,
            "last_updated": native.get("LAST_UPDATED"),
            "_join_keys": {
                "tafs": native.get("TAFS"),
                "agency_identifier": _code(native.get("TRAG"), 3),
                "main_account_code": _code(native.get("TRACCT"), 4),
                "omb_account": native.get("OMB_ACCT"),
            },
            "semantics": "monthly budget execution; not an obligation, award, or opportunity",
        }
        normalized.append(row)
    return normalized, dropped


class OmbSf133Source(DataSource):
    name = SOURCE_NAME
    kind = SourceKind.ENRICHMENT

    def __init__(self, fetcher: Optional[FetchBytes] = None) -> None:
        self._fetcher = fetcher or fetch_bytes

    def search(self, query: SourceQuery) -> list:
        raise NotImplementedError("enrichment-only source; use budget_execution()")

    def healthcheck(self) -> tuple[bool, str]:
        try:
            blob, _headers = self._fetcher(
                _url(AGENCY_SLUGS["dhs"]),
                headers={"Range": "bytes=0-4095"}, timeout=20.0, retries=1,
            )
            if blob.startswith(b"PK"):
                return True, "official FY2026 SF-133 workbook reachable"
            return False, "SF-133 endpoint returned a non-XLSX response"
        except Exception as exc:  # noqa: BLE001
            return False, f"SF-133 workbook unavailable: {exc}"

    def budget_execution(self, query: SourceQuery) -> dict[str, Any]:
        slug, limitations = _agency_slug(query)
        if slug is None:
            return payload(
                source=self.name, mode=MODE, items=[], status="partial",
                limitations=limitations,
            )
        url = _url(slug)
        try:
            blob, headers = self._fetcher(url, timeout=90.0, retries=2)
            matched, dropped = _parse_workbook(blob, query)
        except Exception as exc:  # noqa: BLE001
            return payload(
                source=self.name, mode=MODE, items=[], status="failed",
                limitations=limitations + [f"SF-133 retrieval or parse failed: {exc}"],
            )
        limit = max(1, min(query.limit, 500))
        items = matched[:limit]
        if dropped:
            limitations.append(f"ignored {dropped} rows without a TAFS or line description")
        if len(matched) > limit:
            limitations.append(f"kept first {limit} of {len(matched)} scoped rows")
        limitations.extend([
            "SF-133 is account-level budget execution, not procurement intent or award evidence",
            "raw preparer identifiers F2_USER_ID and F2_USER are deliberately discarded",
            "only explicitly mapped, independently verified FY2026 agency workbook URLs are supported",
        ])
        source_receipt = receipt(
            source=self.name, url=url,
            query={"agencies": query.agencies, "keywords": query.keywords,
                   "limit": query.limit, "fiscal_year": FISCAL_YEAR},
            blob=blob, normalized=items, headers=headers,
            content_kind="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return payload(
            source=self.name, mode=MODE, items=items,
            status="partial" if dropped or len(matched) > limit else "complete",
            limitations=limitations, source_receipt=source_receipt,
            total_matched=len(matched),
            public_detail="Official OMB FY2026 monthly SF-133 workbook",
        )


try:
    register_source(OmbSf133Source)
except ValueError:
    # Integration adds the catalog identity before importing this module.
    pass
