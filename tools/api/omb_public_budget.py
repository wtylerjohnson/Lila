"""OMB Public Budget Database adapter (Budget Authority workbook).

Client dimension: budget context.  The workbook provides account-level
President's Budget history and estimates, keyed by agency/bureau/account,
Treasury agency, CGAC agency, function, and BEA category.  It is planning
evidence, never proof of enacted authority, apportionment, or obligations.

One official landing-page read discovers the versioned workbook; one bounded
binary read retrieves it.  Filtering happens locally because OMB publishes no
query API.  The returned rows are scoped and capped, while provenance names the
annual-bulk-file limitation.
"""

from __future__ import annotations

import hashlib
import io
import re
import time
from datetime import datetime, timezone
from html import unescape
from typing import Any, Callable, Iterable, Optional
from urllib.parse import urljoin

from tools.api import _http
from tools.api._http import get_text
from tools.api.base import (DataSource, SourceKind, SourceQuery,
                            cap_disclosed, register_source)
from tools.api.provenance import make_provenance_envelope

SOURCE_NAME = "omb_public_budget_database"
LANDING_URL = (
    "https://www.whitehouse.gov/omb/information-resources/budget/"
    "supplemental-materials/"
)
WORKBOOK_RE = re.compile(
    r'href=["\']([^"\']*budauth_fy\d{4}\.xlsx(?:\?[^"\']*)?)["\']',
    re.IGNORECASE,
)
BUDGET_YEAR_RE = re.compile(r"budauth_fy(\d{4})\.xlsx", re.IGNORECASE)
MAX_DOWNLOAD_BYTES = 12 * 1024 * 1024
MAX_RETURNED = 100

IDENTITY_HEADERS = (
    "Agency Code", "Agency Name", "Bureau Code", "Bureau Name",
    "Account Code", "Account Name", "Treasury Agency Code",
    "CGAC Agency Code", "Subfunction Code", "Subfunction Title",
    "BEA Category", "On- or Off- Budget",
)


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _workbook_url(html: str) -> Optional[str]:
    match = WORKBOOK_RE.search(html or "")
    return urljoin(LANDING_URL, unescape(match.group(1))) if match else None


def _budget_year(url: str) -> Optional[int]:
    match = BUDGET_YEAR_RE.search(url or "")
    return int(match.group(1)) if match else None


def _get_bytes(url: str, *, timeout: float = 60.0,
               retries: int = 3) -> bytes:
    last_exc: Optional[Exception] = None
    for attempt in range(retries):
        try:
            with _http._client(timeout) as client:
                response = client.get(url)
            if response.status_code in _http._RETRYABLE:
                raise RuntimeError(f"HTTP {response.status_code} fetching workbook")
            if response.status_code >= 400:
                raise RuntimeError(f"HTTP {response.status_code} fetching workbook")
            declared = int(response.headers.get("content-length") or 0)
            if declared > MAX_DOWNLOAD_BYTES or len(response.content) > MAX_DOWNLOAD_BYTES:
                raise ValueError(
                    f"OMB workbook exceeds {MAX_DOWNLOAD_BYTES} byte safety cap")
            return bytes(response.content)
        except Exception as exc:  # noqa: BLE001 - bounded retry
            last_exc = exc
            if attempt < retries - 1:
                time.sleep(0.5 * (2 ** attempt))
    assert last_exc is not None
    raise last_exc


def _grid_from_xlsx(blob: bytes) -> Iterable[tuple[Any, ...]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required for OMB workbook ingestion") from exc
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def _clean(value: Any) -> Optional[str]:
    text = " ".join(str(value or "").split())
    return text or None


def _code(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean(value)


def _amount(value: Any) -> Optional[float | int]:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return int(value) if float(value).is_integer() else float(value)


def _record_id(row: dict) -> str:
    keys = (
        "agency_code", "bureau_code", "account_code", "treasury_agency_code",
        "cgac_agency_code", "subfunction_code", "bea_category", "on_off_budget",
    )
    basis = "|".join(str(row.get(key) or "").casefold() for key in keys)
    return "omb-budget:" + hashlib.sha256(basis.encode("utf-8")).hexdigest()[:20]


def _rows_from_grid(grid: Iterable[tuple[Any, ...]], *,
                    budget_year: Optional[int]) -> tuple[list[dict], int]:
    iterator = iter(grid)
    try:
        headers = [str(value or "").strip() for value in next(iterator)]
    except StopIteration:
        raise ValueError("OMB workbook is empty")
    missing = [name for name in IDENTITY_HEADERS if name not in headers]
    if missing:
        raise ValueError("OMB workbook schema drift; missing: " + ", ".join(missing))
    index = {name: position for position, name in enumerate(headers)}
    year_headers = [
        (position, int(name)) for position, name in enumerate(headers)
        if name.isdigit() and len(name) == 4
    ]
    if not year_headers:
        raise ValueError("OMB workbook schema drift; no fiscal-year columns")
    # Preserve the decision-useful budget window, not 50+ historical columns.
    if budget_year:
        selected_years = [
            pair for pair in year_headers
            if budget_year - 2 <= pair[1] <= budget_year + 4
        ]
    else:
        selected_years = year_headers[-7:]

    rows: list[dict] = []
    dropped = 0
    width = max(index.values()) + 1
    for native in iterator:
        if not native or all(value in (None, "") for value in native):
            continue
        if len(native) < width:
            dropped += 1
            continue
        row = {
            "agency_code": _code(native[index["Agency Code"]]),
            "agency_name": _clean(native[index["Agency Name"]]),
            "bureau_code": _code(native[index["Bureau Code"]]),
            "bureau_name": _clean(native[index["Bureau Name"]]),
            "account_code": _code(native[index["Account Code"]]),
            "account_name": _clean(native[index["Account Name"]]),
            "treasury_agency_code": _code(
                native[index["Treasury Agency Code"]]),
            "cgac_agency_code": _code(native[index["CGAC Agency Code"]]),
            "subfunction_code": _code(native[index["Subfunction Code"]]),
            "subfunction_title": _clean(native[index["Subfunction Title"]]),
            "bea_category": _clean(native[index["BEA Category"]]),
            "on_off_budget": _clean(native[index["On- or Off- Budget"]]),
            "year_amounts": {
                str(year): amount
                for position, year in selected_years
                if (amount := _amount(
                    native[position] if position < len(native) else None)) is not None
            },
            "amount_unit": "thousands_usd",
        }
        if not row["agency_name"] or not row["account_name"]:
            dropped += 1
            continue
        row["record_id"] = _record_id(row)
        rows.append(row)
    return rows, dropped


def _terms(query: SourceQuery) -> list[str]:
    values = list(query.agencies) + list(query.keywords)
    seen: set[str] = set()
    result = []
    for value in values:
        term = str(value or "").strip().strip('"').casefold()
        if term and term not in seen:
            seen.add(term)
            result.append(term)
    return result


def _matches(row: dict, terms: list[str]) -> list[str]:
    haystack = " ".join(str(row.get(key) or "") for key in (
        "agency_code", "agency_name", "bureau_code", "bureau_name",
        "account_code", "account_name", "treasury_agency_code",
        "cgac_agency_code", "subfunction_code", "subfunction_title",
    )).casefold()
    return [term for term in terms if term in haystack]


class OmbPublicBudgetSource(DataSource):
    name = SOURCE_NAME
    kind = SourceKind.ENRICHMENT

    def __init__(
        self,
        *,
        fetch_text: Optional[Callable[..., str]] = None,
        fetch_bytes: Optional[Callable[..., bytes]] = None,
        parse_workbook: Optional[Callable[..., Iterable[tuple[Any, ...]]]] = None,
    ) -> None:
        self._fetch_text = fetch_text or get_text
        self._fetch_bytes = fetch_bytes or _get_bytes
        self._parse_workbook = parse_workbook or _grid_from_xlsx

    def search(self, query: SourceQuery) -> list:
        raise NotImplementedError(
            "enrichment-only source; use budget_authority()")

    def healthcheck(self) -> tuple[bool, str]:
        try:
            html = self._fetch_text(LANDING_URL, timeout=8.0, retries=1)
        except Exception as exc:  # noqa: BLE001
            return False, f"OMB supplemental-materials page unreachable: {exc}"
        url = _workbook_url(html)
        return ((True, f"current Budget Authority workbook indexed: {url}")
                if url else (False, "Budget Authority workbook link not found"))

    def budget_authority(self, query: SourceQuery) -> dict[str, Any]:
        terms = _terms(query)
        if not terms:
            return self._failure(
                "unscoped query refused; pass agency or account/program keywords",
                status="partial",
            )
        retrieved_at = _utc_now()
        try:
            html = self._fetch_text(LANDING_URL, timeout=20.0, retries=2)
            workbook_url = _workbook_url(html)
            if not workbook_url:
                raise ValueError("Budget Authority workbook link not found")
            year = _budget_year(workbook_url)
            blob = self._fetch_bytes(workbook_url, timeout=60.0, retries=2)
            all_rows, dropped = _rows_from_grid(
                self._parse_workbook(blob), budget_year=year)
        except Exception as exc:  # noqa: BLE001 - adapter failure isolation
            return self._failure(str(exc), retrieved_at=retrieved_at)

        matched: list[dict] = []
        for row in all_rows:
            hits = _matches(row, terms)
            if hits:
                matched.append({**row, "matched_terms": hits,
                                "url": workbook_url})
        matched.sort(key=lambda row: (
            row.get("agency_name") or "", row.get("bureau_name") or "",
            row.get("account_name") or "", row["record_id"],
        ))
        limit = max(1, min(query.limit, MAX_RETURNED))
        payload = cap_disclosed(matched, limit, key="accounts")
        limitations = [
            "President's Budget proposal; not enacted authority, apportionment, or obligations",
            "annual bulk workbook is locally filtered because OMB exposes no query API",
            "amounts are thousands of dollars as published by OMB",
        ]
        status = "partial" if dropped else "complete"
        if dropped:
            limitations.append(f"{dropped} malformed workbook rows dropped")
        payload.update({
            "budget_year": year,
            "workbook_url": workbook_url,
            "landing_url": LANDING_URL,
            "rows_scanned": len(all_rows) + dropped,
            "rows_dropped": dropped,
            "_provenance": make_provenance_envelope(
                self.name,
                status=status,
                mode="live_version_discovery_bulk_xlsx",
                retrieval_mode="live",
                retrieved_at=retrieved_at,
                data_as_of=str(year) if year else None,
                record_count=len(payload["accounts"]),
                limitations=limitations,
            ).model_dump(mode="json") | {
                "retrieval_url": workbook_url,
                "retrieval_query": query.model_dump(mode="json"),
                "raw_content_sha256": hashlib.sha256(blob).hexdigest(),
                "raw_content_kind": (
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                "normalization_version": "omb_public_budget.v1.2026-08-08",
            },
        })
        return payload

    def _failure(self, error: str, *, status: str = "failed",
                 retrieved_at: Optional[datetime] = None) -> dict:
        return {
            "accounts": [],
            "total_matched": 0,
            "error": error,
            "landing_url": LANDING_URL,
            "_provenance": make_provenance_envelope(
                self.name,
                status=status,
                mode="live_version_discovery_bulk_xlsx",
                retrieval_mode="live",
                retrieved_at=retrieved_at or _utc_now(),
                record_count=0,
                limitations=[error],
            ).model_dump(mode="json"),
        }


try:
    register_source(OmbPublicBudgetSource)
except ValueError:
    pass
