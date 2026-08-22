"""NASA Consolidated Agency-wide Acquisition Forecast (NAF) adapter.

NASA publishes an official agency-wide forecast workbook with requirement
titles, buying centers, NAICS and PSC codes, stated value bands, forecast
quarters, acquisition lifecycle state, and named NASA contacts.  These rows
are planning intent, not solicitations.  This adapter therefore emits only
``ForecastRecord`` objects and its discovery ``search`` method is deliberately
empty.

Official workbook:
    https://www.hq.nasa.gov/office/procurement/forecast/AcqForecastNew.xlsx
Official landing page:
    https://www.nasa.gov/osbp/acquisition-forecast/

Rows marked Awarded or Withdrawn remain source history and are excluded from
the forward forecast set.  NASA's ``SourceID`` is the stable identifier basis;
when it is absent, a deterministic hash of the agency-published requirement
fields is used instead.
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import tempfile
import time
from datetime import date, datetime, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Callable, Mapping, Optional

from openpyxl import load_workbook

from agents.schemas import ForecastRecord, RawOpportunity
from tools.api import _http
from tools.api.base import (
    DataSource,
    SourceKind,
    SourceQuery,
    register_source,
)
from tools.api.provenance import make_provenance_envelope
from tools.toggles import is_enabled

logger = logging.getLogger(__name__)

SOURCE_NAME = "nasa_naf"
WORKBOOK_URL = os.environ.get(
    "LILA_NASA_NAF_URL",
    "https://www.hq.nasa.gov/office/procurement/forecast/AcqForecastNew.xlsx",
)
LANDING_URL = "https://www.nasa.gov/osbp/acquisition-forecast/"
SHEET_NAME = "Forecast"
HEADERS = {
    "User-Agent": (
        "GTM-Group-LILA/1.0 (federal capture research; "
        "william.tyler.johnson@gmail.com)"
    ),
    "Accept": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,"
        "application/octet-stream;q=0.9"
    ),
}
_DEFAULT_CACHE = (
    Path(__file__).resolve().parents[3] / "data" / "cache" / "forecasts"
)
_NULL_TEXT = {"", "n/a", "none", "null"}
_INACTIVE_STATES = {"awarded", "withdrawn"}
_ID_FALLBACK_FIELDS = (
    "TitleOfRequirement",
    "BuyingOffice",
    "NAICS",
    "FYofSolOrNOFORelease",
    "QtrSolOrNOFORelease",
)

FetchResult = tuple[bytes, dict[str, Any]]
FetchFn = Callable[[str], FetchResult]
ProbeFn = Callable[[str], dict[str, Any]]


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _today() -> date:
    return _now().date()


def _cache_dir() -> Path:
    return Path(
        os.environ.get("LILA_FORECAST_CACHE_DIR", str(_DEFAULT_CACHE))
    )


def _get_xlsx(
    url: str,
    *,
    timeout: float = 30.0,
    retries: int = 3,
) -> FetchResult:
    """Fetch workbook bytes while retaining source HTTP metadata."""

    last_error: Optional[Exception] = None
    for attempt in range(retries):
        try:
            with _http._client(timeout) as client:
                response = client.get(url, headers=HEADERS)
            if response.status_code in _http._RETRYABLE:
                raise RuntimeError(
                    f"retryable HTTP {response.status_code} from {url}"
                )
            response.raise_for_status()
            return response.content, {
                "url": str(response.url),
                "last_modified": response.headers.get("Last-Modified"),
                "etag": response.headers.get("ETag"),
                "content_length": response.headers.get("Content-Length"),
            }
        except Exception as exc:  # noqa: BLE001 - bounded HTTP retry loop
            last_error = exc
            if attempt < retries - 1:
                time.sleep(0.5 * (2**attempt))
    assert last_error is not None
    raise last_error


def _head_xlsx(url: str, *, timeout: float = 8.0) -> dict[str, Any]:
    """Probe the real workbook URL without downloading the body."""

    with _http._client(timeout) as client:
        response = client.request("HEAD", url, headers=HEADERS)
    response.raise_for_status()
    return {
        "status_code": response.status_code,
        "last_modified": response.headers.get("Last-Modified"),
        "etag": response.headers.get("ETag"),
        "content_length": response.headers.get("Content-Length"),
    }


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.casefold() in _NULL_TEXT:
        return None
    return text


def _canonical_id_part(value: Any) -> str:
    """Normalize fallback-ID fields without changing published display text."""

    return " ".join((_text(value) or "").split()).casefold()


def _stable_id(row: Mapping[str, Any]) -> str:
    """Preserve NASA's published SourceID; hash only rows without one."""

    source_id = _text(row.get("SourceID"))
    if source_id:
        return source_id
    parts = [_canonical_id_part(row.get(key)) for key in _ID_FALLBACK_FIELDS]
    basis = "row:" + "|".join(parts)
    digest = hashlib.sha256(basis.encode("utf-8")).hexdigest()[:16]
    return f"nasa-naf-fallback-{digest}"


def _naics(value: Any) -> tuple[Optional[str], Optional[str]]:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _text(value)
    if not text:
        return None, None
    parts = text.split(" - ", 1)
    code = parts[0].strip()
    label = parts[1].strip() if len(parts) > 1 else None
    return (code if code.isdigit() and len(code) == 6 else None), label


def _timing(quarter: Any, fiscal_year: Any) -> Optional[str]:
    qtr = _text(quarter)
    year = _text(fiscal_year)
    if qtr and year:
        year = year[2:] if year.upper().startswith("FY") else year
        return f"{qtr} FY{year}"
    if year:
        return year if year.upper().startswith("FY") else f"FY{year}"
    return qtr


def _contact(name: Any, email: Any) -> Optional[str]:
    person = _text(name)
    address = _text(email)
    return " ".join(
        part
        for part in (person, f"<{address}>" if address else None)
        if part
    ) or None


def _http_date_to_iso(value: Any) -> Optional[str]:
    text = _text(value)
    if not text:
        return None
    try:
        parsed = parsedate_to_datetime(text)
    except (IndexError, TypeError, ValueError):
        return text
    if parsed is None:
        return text
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc).isoformat()


def rows_from_xlsx(payload: bytes) -> list[dict[str, Any]]:
    """Parse the official workbook into full-width, header-keyed rows."""

    workbook = load_workbook(
        io.BytesIO(payload), read_only=True, data_only=True
    )
    try:
        if not workbook.sheetnames:
            raise ValueError("NASA forecast workbook contains no sheets")
        sheet = (
            workbook[SHEET_NAME]
            if SHEET_NAME in workbook.sheetnames
            else workbook[workbook.sheetnames[0]]
        )
        rows = sheet.iter_rows(values_only=True)
        header_cells = next(rows, None)
        if not header_cells:
            raise ValueError("NASA forecast workbook contains no header row")
        headers = [str(cell).strip() if cell is not None else "" for cell in header_cells]
        if "TitleOfRequirement" not in headers:
            raise ValueError("NASA forecast workbook is missing TitleOfRequirement")
        output: list[dict[str, Any]] = []
        for cells in rows:
            if cells is None or all(cell in (None, "") for cell in cells):
                continue
            row: dict[str, Any] = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = cells[index] if index < len(cells) else None
                if isinstance(value, (date, datetime)):
                    value = value.isoformat()
                row[header] = value
            output.append(row)
        return output
    finally:
        workbook.close()


def _is_inactive(row: Mapping[str, Any]) -> bool:
    status = (_text(row.get("AcquisitionStatus")) or "").casefold()
    lifecycle = (_text(row.get("AwardedOrWithdrawn")) or "").casefold()
    return status in _INACTIVE_STATES or lifecycle in _INACTIVE_STATES


def map_row(
    row: Mapping[str, Any],
    *,
    retrieved_at: datetime,
    data_as_of: Optional[str],
) -> Optional[ForecastRecord]:
    """Map one NASA-published row without promoting it to a solicitation."""

    title = _text(row.get("TitleOfRequirement"))
    if not title:
        return None
    naics_code, embedded_label = _naics(row.get("NAICS"))
    naics_label = _text(row.get("NAICS Description")) or embedded_label
    solicitation_fy = _text(row.get("FYofSolOrNOFORelease"))
    award_fy = _text(row.get("AnticipatedFYAward"))
    status = (
        _text(row.get("AcquisitionStatus"))
        or _text(row.get("AcquisitionPhase"))
        or "Agency acquisition forecast"
    )
    vehicle = _text(row.get("Type of Award/Contract Vehicle"))
    contract_type = _text(row.get("ContractType"))
    return ForecastRecord(
        source=SOURCE_NAME,
        source_id=_stable_id(row),
        agency="NASA",
        component=_text(row.get("BuyingOffice")),
        title=title,
        description=(
            _text(row.get("Summary")) or _text(row.get("Description"))
        ),
        naics_code=naics_code,
        naics_label=naics_label,
        psc=_text(row.get("PSC Code")),
        estimated_value_range=_text(row.get("EstimatedContractValue")),
        anticipated_solicitation=_timing(
            row.get("QtrSolOrNOFORelease"), solicitation_fy
        ),
        anticipated_award=_timing(
            row.get("Anticipated Qtr of Award"), award_fy
        ),
        fiscal_year=solicitation_fy or award_fy,
        award_type=contract_type or vehicle,
        set_aside=_text(row.get("SetAsideType")),
        small_business_poc=_contact(
            row.get("SmallBusinessSpecialistPOC"),
            row.get("SmallBusinessSpecialistEmail"),
        ),
        url=WORKBOOK_URL,
        retrieved_at=retrieved_at,
        data_as_of=data_as_of,
        forecast_status=status,
        incumbent_stated=None,
        source_fields=dict(row),
    )


def _atomic_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f"{path.name}.", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f"{path.name}.", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, sort_keys=True)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass


def _read_meta(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def _parse_retrieved(value: Any, *, fallback_path: Optional[Path] = None) -> datetime:
    text = _text(value)
    if text:
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=timezone.utc)
            return parsed.astimezone(timezone.utc)
        except ValueError:
            pass
    if fallback_path is not None:
        return datetime.fromtimestamp(
            fallback_path.stat().st_mtime, tz=timezone.utc
        )
    return _now()


class NasaNafSource(DataSource):
    """Official NASA forecast collector."""

    name = SOURCE_NAME
    kind = SourceKind.ENRICHMENT
    forecast_agency = "NASA"

    def __init__(
        self,
        *,
        fetch: Optional[FetchFn] = None,
        probe: Optional[ProbeFn] = None,
    ) -> None:
        self.enabled = is_enabled(self.name, True)
        self.last_provenance: dict[str, Any] = {}
        self._fetch = fetch or _get_xlsx
        self._probe = probe or _head_xlsx

    def search(self, query: SourceQuery) -> list[RawOpportunity]:
        """Forecast intent never enters the active-opportunity population."""

        return []

    def healthcheck(self) -> tuple[bool, str]:
        try:
            metadata = self._probe(WORKBOOK_URL) or {}
        except Exception as exc:  # noqa: BLE001 - health probes report failures
            return False, f"official NASA forecast workbook unavailable: {exc}"
        details = []
        if metadata.get("content_length"):
            details.append(f"{metadata['content_length']} bytes")
        if metadata.get("last_modified"):
            details.append(f"last modified {metadata['last_modified']}")
        suffix = f" ({', '.join(details)})" if details else ""
        return True, f"official NASA forecast workbook reachable{suffix}"

    def _provenance(
        self,
        *,
        status: str,
        mode: str,
        retrieval_mode: str,
        retrieved_at: datetime,
        data_as_of: Optional[str],
        record_count: int,
        total_rows: int,
        withdrawn_rows: int,
        malformed_rows: int,
        fallback: bool = False,
        stale: bool = False,
        private_error: Optional[str] = None,
    ) -> None:
        limitations = [
            "Agency acquisition forecast; planning intent, never an active "
            "solicitation or response deadline.",
            "The workbook publishes no incumbent name or SAM.gov notice ID "
            "for pre-solicitation rows.",
        ]
        if withdrawn_rows:
            limitations.append(
                f"{withdrawn_rows} rows marked Awarded or Withdrawn excluded "
                "from the forward forecast set."
            )
        if malformed_rows:
            limitations.append(
                f"{malformed_rows} source rows omitted because no requirement "
                "title could be normalized."
            )
        if stale:
            limitations.append(
                "Live official workbook refresh failed; using the newest "
                "parseable official workbook cache."
            )
        envelope = make_provenance_envelope(
            self.name,
            status=status,  # type: ignore[arg-type]
            mode=mode,
            retrieval_mode=retrieval_mode,  # type: ignore[arg-type]
            fallback=fallback,
            stale=stale,
            retrieved_at=retrieved_at,
            data_as_of=data_as_of,
            record_count=record_count,
            attempts=[{
                "source": self.name,
                "status": "success" if status == "complete" else status,
                "count": record_count,
                "retrieved_at": retrieved_at,
                "data_as_of": data_as_of,
                "error": private_error,
            }],
            limitations=limitations,
            public_detail=(
                f"Official NASA acquisition forecast workbook; "
                f"{record_count} forward records"
            ),
        ).model_dump(mode="json")
        self.last_provenance = {
            **envelope,
            "complete": status == "complete",
            "records": record_count,
            "total_available": total_rows,
            "inactive_rows_excluded": withdrawn_rows,
            "malformed_rows": malformed_rows,
            "source_url": WORKBOOK_URL,
            "landing_url": LANDING_URL,
            "classification": "agency forecast; never an active notice",
        }
        if private_error:
            self.last_provenance["private_error"] = private_error

    def _load_snapshot(self) -> tuple[list[dict[str, Any]], dict[str, Any], str]:
        day = _today().isoformat()
        cache = _cache_dir() / f"nasa_naf_{day}.xlsx"
        meta_path = _cache_dir() / f"nasa_naf_{day}.meta.json"
        if cache.exists():
            try:
                payload = cache.read_bytes()
                rows = rows_from_xlsx(payload)
                metadata = _read_meta(meta_path)
                metadata.setdefault(
                    "retrieved_at",
                    datetime.fromtimestamp(
                        cache.stat().st_mtime, tz=timezone.utc
                    ).isoformat(),
                )
                return rows, metadata, "official_daily_cache"
            except (OSError, ValueError):
                pass

        try:
            payload, fetched = self._fetch(WORKBOOK_URL)
            rows = rows_from_xlsx(payload)
            metadata = dict(fetched or {})
            metadata["retrieved_at"] = _now().isoformat()
            _atomic_bytes(cache, payload)
            _atomic_json(meta_path, metadata)
            return rows, metadata, "live_official_workbook"
        except Exception as live_error:  # noqa: BLE001 - named stale fallback
            candidates = sorted(
                _cache_dir().glob("nasa_naf_????-??-??.xlsx"), reverse=True
            )
            for candidate in candidates:
                if candidate == cache:
                    continue
                try:
                    rows = rows_from_xlsx(candidate.read_bytes())
                except (OSError, ValueError):
                    continue
                metadata = _read_meta(candidate.with_suffix(".meta.json"))
                metadata.setdefault(
                    "retrieved_at",
                    datetime.fromtimestamp(
                        candidate.stat().st_mtime, tz=timezone.utc
                    ).isoformat(),
                )
                metadata["private_error"] = str(live_error)
                return rows, metadata, "stale_official_cache"
            raise

    def forecasts(
        self, query: Optional[SourceQuery] = None
    ) -> list[ForecastRecord]:
        """Return the complete forward NASA forecast as typed program signals."""

        if not self.enabled:
            now = _now()
            self._provenance(
                status="failed",
                mode="not_run",
                retrieval_mode="unknown",
                retrieved_at=now,
                data_as_of=None,
                record_count=0,
                total_rows=0,
                withdrawn_rows=0,
                malformed_rows=0,
            )
            return []
        try:
            rows, metadata, mode = self._load_snapshot()
        except Exception as exc:  # noqa: BLE001 - one lane cannot sink the sweep
            now = _now()
            self._provenance(
                status="failed",
                mode="live_official_workbook",
                retrieval_mode="live",
                retrieved_at=now,
                data_as_of=None,
                record_count=0,
                total_rows=0,
                withdrawn_rows=0,
                malformed_rows=0,
                private_error=str(exc),
            )
            logger.warning(
                "nasa_naf: official workbook pull failed; returning no rows: %s",
                exc,
            )
            return []

        retrieved_at = _parse_retrieved(metadata.get("retrieved_at"))
        data_as_of = _http_date_to_iso(metadata.get("last_modified"))
        withdrawn_rows = 0
        malformed_rows = 0
        output: list[ForecastRecord] = []
        for row in rows:
            if _is_inactive(row):
                withdrawn_rows += 1
                continue
            try:
                mapped = map_row(
                    row,
                    retrieved_at=retrieved_at,
                    data_as_of=data_as_of,
                )
            except Exception:  # noqa: BLE001 - one malformed row is isolated
                malformed_rows += 1
                continue
            if mapped is None:
                malformed_rows += 1
                continue
            output.append(mapped)

        status = "partial" if malformed_rows else "complete"
        is_stale = mode == "stale_official_cache"
        retrieval_mode = "official-cache" if "cache" in mode else "live"
        self._provenance(
            status=status,
            mode=mode,
            retrieval_mode=retrieval_mode,
            retrieved_at=retrieved_at,
            data_as_of=data_as_of,
            record_count=len(output),
            total_rows=len(rows),
            withdrawn_rows=withdrawn_rows,
            malformed_rows=malformed_rows,
            fallback=is_stale,
            stale=is_stale,
            private_error=metadata.get("private_error"),
        )
        return output


def _register_when_cataloged() -> None:
    """Allow the parent wiring change to activate the adapter atomically."""

    from tools.api.source_catalog import SOURCE_BY_ADAPTER

    if NasaNafSource.name in SOURCE_BY_ADAPTER:
        register_source(NasaNafSource)
    else:
        logger.info(
            "nasa_naf registration deferred until source-catalog wiring"
        )


_register_when_cataloged()
