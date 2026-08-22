"""Public Army OSBP acquisition-forecast workbook adapter.

The Army Office of Small Business Programs publishes its acquisition forecast
as an official XLSX workbook.  This adapter downloads and validates that source
document, preserves the agency's identifiers, value bands, dates, and workbook
fields, and emits only :class:`ForecastRecord` objects.  It never emits a
``RawOpportunity`` and therefore cannot enter the active-notice flow.

The workbook URL is versioned by the Army content service.  Operators may point
``LILA_ARMY_FORECAST_URL`` at a newer official workbook without changing code.
Each successful document is cached once per UTC day; a failed live refresh may
fall back to the newest parseable official workbook cache with that limitation
named in provenance.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from agents.schemas import ForecastRecord
from tools.api._http import get_bytes
from tools.api.base import DataSource, SourceKind, SourceQuery, register_source
from tools.toggles import is_enabled

WORKBOOK_URL = os.environ.get(
    "LILA_ARMY_FORECAST_URL",
    "https://api.army.mil/e2/c/downloads/2026/06/18/565277f2/"
    "2026-osbp-acq-forecast-jul-dec.xlsx",
)
LANDING_URL = "https://www.army.mil/osbp"
SHEET_NAME = "Preaward_OSBP"
HEADERS = {
    "User-Agent": (
        "GTM-Group-LILA/1.0 (federal capture research; "
        "william.tyler.johnson@gmail.com)"
    ),
    "Accept": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,"
        "application/octet-stream;q=0.9"
    ),
}
MAX_WORKBOOK_BYTES = int(
    os.environ.get("LILA_ARMY_FORECAST_MAX_BYTES", str(15 * 1024 * 1024))
)
_DEFAULT_CACHE = Path(__file__).resolve().parents[3] / "data" / "cache" / "forecasts"
_MONTHS = {
    name: number
    for number, name in enumerate(
        (
            "january",
            "february",
            "march",
            "april",
            "may",
            "june",
            "july",
            "august",
            "september",
            "october",
            "november",
            "december",
        ),
        1,
    )
}


def _today() -> date:
    return datetime.now(timezone.utc).date()


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _cache_dir() -> Path:
    return Path(
        os.environ.get("LILA_FORECAST_CACHE_DIR", str(_DEFAULT_CACHE))
    )


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


_FIELD_BY_HEADER = {
    "pan solicitation or contract number": "source_id",
    "if follow on provide current contract number": "predecessor_contract_id",
    "description of requirement": "title",
    "consolidation anticipated": "consolidation",
    "bundling anticipated": "bundling",
    "forecasted contract value": "estimated_value_range",
    "forecasted psc": "psc",
    "anticipated naics": "naics_code",
    "anticipated type of set aside or unknown": "set_aside",
    "procact": "procurement_action_type",
    "indefinite delivery vehicle type for task delivery orders or unknown": (
        "idv_type"
    ),
    "palt": "palt",
    "anticipated contract type or unknown": "contract_type",
    "anticipated solicitation date": "anticipated_solicitation",
    "anticipated solicitation closing date": "anticipated_solicitation_close",
    "forecasted award date": "anticipated_award",
    "period of performance": "period_of_performance",
    "contracting center": "contracting_center",
    "contracting office": "contracting_office",
    "command": "command",
    "pm directorate": "directorate",
    "assigned small business office": "small_business_office",
    "assigned small business office email": "small_business_office_email",
}
_REQUIRED_FIELDS = {"source_id", "title", "estimated_value_range"}


def _text(value: Any) -> Optional[str]:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = " ".join(str(value).split()).strip()
    return text or None


def _date_text(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _text(value)


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _document_dates(title: str) -> tuple[Optional[str], Optional[str]]:
    fiscal_year = None
    fiscal_match = re.search(r"\bFY\s*(\d{2,4})\b", title, re.IGNORECASE)
    if fiscal_match:
        number = int(fiscal_match.group(1))
        fiscal_year = str(2000 + number if number < 100 else number)
    data_as_of = None
    month_match = re.search(
        r"\b(" + "|".join(_MONTHS) + r")\s+(\d{4})\b",
        title.casefold(),
    )
    if month_match:
        data_as_of = (
            f"{month_match.group(2)}-{_MONTHS[month_match.group(1)]:02d}"
        )
    return fiscal_year, data_as_of


def _row_rank(row: dict[str, Any]) -> tuple[int, str]:
    """Order duplicate official rows without depending on workbook position."""

    populated = sum(
        value not in (None, "")
        for key, value in row.items()
        if not key.startswith("_")
    )
    canonical = json.dumps(
        {
            key: _json_value(value)
            for key, value in row.items()
            if not key.startswith("_")
        },
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return populated, canonical


def _stable_unique(values: list[Any]) -> list[Any]:
    """Return JSON-safe published values in deterministic content order."""

    by_key: dict[str, Any] = {}
    for value in values:
        if value in (None, ""):
            continue
        safe = _json_value(value)
        key = json.dumps(safe, sort_keys=True, separators=(",", ":"), default=str)
        by_key.setdefault(key, safe)
    return [by_key[key] for key in sorted(by_key)]


def _merge_same_id_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Merge repeated Army identifiers while retaining every stated value.

    Army's workbook repeats some PAN/solicitation identifiers.  The repeated
    lines are often complementary (one line supplies NAICS while another does
    not), and a small number carry conflicting stated values.  Selecting one
    row discards official evidence.  The merged scalar view therefore starts
    with the most complete row and fills its blanks, while the complete set of
    distinct published values and source rows remains attached for audit and
    matching.
    """

    ranked = sorted(rows, key=_row_rank, reverse=True)
    merged = dict(ranked[0])
    merged["_source_fields"] = dict(ranked[0].get("_source_fields") or {})
    public_keys = sorted(
        {
            key
            for row in rows
            for key in row
            if not key.startswith("_")
        }
    )
    published_values: dict[str, list[Any]] = {}
    for key in public_keys:
        values = _stable_unique([row.get(key) for row in rows])
        if not values:
            continue
        if merged.get(key) in (None, ""):
            merged[key] = values[0]
        if len(values) > 1:
            published_values[key] = values

    source_rows = [
        {
            str(key): _json_value(value)
            for key, value in (row.get("_source_fields") or {}).items()
        }
        for row in rows
    ]
    source_rows.sort(
        key=lambda row: json.dumps(
            row, sort_keys=True, separators=(",", ":"), default=str
        )
    )
    for source_row in source_rows:
        for header, value in source_row.items():
            if merged["_source_fields"].get(header) in (None, ""):
                merged["_source_fields"][header] = value
    merged["_merged_source_rows"] = source_rows
    merged["_published_values"] = published_values
    merged["_source_row_numbers"] = sorted(
        int(row.get("_row_number") or 0) for row in rows
    )
    merged["_row_number"] = min(merged["_source_row_numbers"])
    return merged


_FALLBACK_ID_FIELDS = (
    "predecessor_contract_id",
    "title",
    "contracting_center",
    "contracting_office",
    "command",
    "directorate",
    "procurement_action_type",
    "psc",
    "naics_code",
)


def _fallback_source_id(row: dict[str, Any]) -> str:
    """Derive a row-order-independent ID when Army publishes no PAN/contract ID."""

    identity = {
        key: (_text(row.get(key)) or "").casefold()
        for key in _FALLBACK_ID_FIELDS
    }
    if not any(identity.values()):
        identity = {
            _normalize_header(key): (_text(value) or "").casefold()
            for key, value in (row.get("_source_fields") or {}).items()
            if _text(value)
        }
    basis = json.dumps(identity, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(basis.encode("utf-8")).hexdigest()[:16]
    return f"army-osbp-fallback-{digest}"


def parse_workbook(payload: bytes) -> dict[str, Any]:
    """Parse and validate one official Army forecast workbook payload."""

    if not payload:
        raise ValueError("Army forecast workbook is empty")
    if len(payload) > MAX_WORKBOOK_BYTES:
        raise ValueError(
            f"Army forecast workbook exceeds {MAX_WORKBOOK_BYTES} bytes"
        )
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - normalize parser-specific failures
        raise ValueError(f"Army forecast workbook is not a valid XLSX: {exc}") from exc

    candidates = list(workbook.worksheets)
    candidates.sort(key=lambda sheet: sheet.title != SHEET_NAME)
    selected = None
    header_row = None
    header_keys: list[Optional[str]] = []
    raw_headers: list[str] = []
    for sheet in candidates:
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True),
            1,
        ):
            keys = [_FIELD_BY_HEADER.get(_normalize_header(value)) for value in values]
            if _REQUIRED_FIELDS.issubset({key for key in keys if key}):
                selected = sheet
                header_row = row_number
                header_keys = keys
                raw_headers = [str(value or "") for value in values]
                break
        if selected is not None:
            break
    if selected is None or header_row is None:
        raise ValueError(
            "Army forecast workbook has no recognized identifier/title/value header row"
        )

    source_title = _text(selected.cell(row=1, column=1).value) or selected.title
    fiscal_year, data_as_of = _document_dates(source_title)
    source_rows = 0
    missing_identifier_rows = 0
    duplicate_source_id_rows = 0
    duplicate_ids: set[str] = set()
    by_id: dict[str, list[dict[str, Any]]] = {}
    for row_number, values in enumerate(
        selected.iter_rows(min_row=header_row + 1, values_only=True),
        header_row + 1,
    ):
        if all(value in (None, "") for value in values):
            continue
        source_rows += 1
        row: dict[str, Any] = {"_row_number": row_number, "_source_fields": {}}
        for index, value in enumerate(values):
            if index >= len(raw_headers) or value in (None, ""):
                continue
            header = raw_headers[index]
            if header:
                row["_source_fields"][header] = _json_value(value)
            key = header_keys[index] if index < len(header_keys) else None
            if key:
                row[key] = value
        source_id = _text(row.get("source_id"))
        if not source_id:
            missing_identifier_rows += 1
            source_id = _fallback_source_id(row)
            row["_generated_source_id"] = True
        row["source_id"] = source_id
        same_id = by_id.setdefault(source_id, [])
        if same_id:
            duplicate_source_id_rows += 1
            duplicate_ids.add(source_id)
        same_id.append(row)

    rows = sorted(
        (_merge_same_id_rows(group) for group in by_id.values()),
        key=lambda row: int(row["_row_number"]),
    )
    if not rows:
        raise ValueError("Army forecast workbook contains no identified records")
    return {
        "source_title": source_title,
        "sheet": selected.title,
        "header_row": header_row,
        "headers": raw_headers,
        "fiscal_year": fiscal_year,
        "data_as_of": data_as_of,
        "source_rows": source_rows,
        "missing_identifier_rows": missing_identifier_rows,
        "generated_identifier_rows": missing_identifier_rows,
        "duplicate_source_id_rows": duplicate_source_id_rows,
        "duplicate_source_ids": sorted(duplicate_ids),
        "rows": rows,
        "sha256": hashlib.sha256(payload).hexdigest(),
    }


def _description(row: dict[str, Any]) -> Optional[str]:
    parts = []
    for label, key in (
        ("Command", "command"),
        ("PM / Directorate", "directorate"),
        ("Contracting office", "contracting_office"),
        ("Contracting center", "contracting_center"),
        ("Procurement action", "procurement_action_type"),
        ("PALT", "palt"),
        ("Period of performance", "period_of_performance"),
    ):
        value = _text(row.get(key))
        if value:
            parts.append(f"{label}: {value}")
    return "; ".join(parts) or None


def map_row(
    row: dict[str, Any],
    *,
    fiscal_year: Optional[str],
    data_as_of: Optional[str],
    retrieved_at: datetime,
) -> ForecastRecord:
    """Map one validated workbook row without promoting it to an opportunity."""

    title = _text(row.get("title")) or ""
    naics = _text(row.get("naics_code"))
    office = _text(row.get("small_business_office"))
    email = _text(row.get("small_business_office_email"))
    small_business_poc = " ".join(
        part
        for part in (office, f"<{email}>" if email else None)
        if part
    ) or None
    award_type = (
        _text(row.get("contract_type"))
        or _text(row.get("idv_type"))
        or _text(row.get("procurement_action_type"))
    )
    return ForecastRecord(
        source="army_acquisition_forecast",
        source_id=str(row["source_id"]),
        agency="Department of the Army",
        component=(
            _text(row.get("command"))
            or _text(row.get("contracting_office"))
            or _text(row.get("contracting_center"))
        ),
        title=title,
        description=_description(row),
        naics_code=naics if naics and re.fullmatch(r"\d{6}", naics) else None,
        psc=_text(row.get("psc")),
        estimated_value_range=_text(row.get("estimated_value_range")),
        anticipated_solicitation=_date_text(row.get("anticipated_solicitation")),
        anticipated_solicitation_close=_date_text(
            row.get("anticipated_solicitation_close")
        ),
        anticipated_award=_date_text(row.get("anticipated_award")),
        fiscal_year=fiscal_year,
        award_type=award_type,
        set_aside=_text(row.get("set_aside")),
        small_business_poc=small_business_poc,
        predecessor_contract_id=_text(row.get("predecessor_contract_id")),
        url=WORKBOOK_URL,
        retrieved_at=retrieved_at,
        data_as_of=data_as_of,
        forecast_status="Agency acquisition forecast",
        source_fields={
            **dict(row.get("_source_fields") or {}),
            **(
                {
                    "_merged_source_rows": list(row["_merged_source_rows"]),
                    "_published_values": dict(row["_published_values"]),
                }
                if row.get("_merged_source_rows")
                and len(row["_merged_source_rows"]) > 1
                else {}
            ),
        },
    )


def _atomic_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f"{path.name}.", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    finally:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass


def _retrieved_at(path: Path) -> datetime:
    return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)


def _latest_valid_cache(exclude: set[Path]) -> Optional[tuple[Path, dict[str, Any]]]:
    paths = sorted(
        _cache_dir().glob("army_acquisition_forecast_????-??-??.xlsx"),
        reverse=True,
    )
    for path in paths:
        if path in exclude:
            continue
        try:
            return path, parse_workbook(path.read_bytes())
        except (OSError, ValueError):
            continue
    return None


@register_source
class ArmyAcquisitionForecastSource(DataSource):
    name = "army_acquisition_forecast"
    kind = SourceKind.DISCOVERY
    forecast_agency = "Department of the Army"

    def __init__(self) -> None:
        self.enabled = is_enabled(self.name, True)
        self.last_provenance: dict[str, Any] = {}

    def search(self, query: SourceQuery) -> list:
        """Army forecast rows never enter the live-opportunity flow."""

        return []

    def _provenance(
        self,
        snapshot: dict[str, Any],
        *,
        mode: str,
        retrieved_at: datetime,
        private_error: Optional[str] = None,
    ) -> None:
        missing = int(snapshot["missing_identifier_rows"])
        duplicate_rows = int(snapshot["duplicate_source_id_rows"])
        limitations = []
        if missing:
            limitations.append(
                f"{missing} source rows had no published identifier and received "
                "deterministic fallback IDs derived from stable published fields"
            )
        if duplicate_rows:
            limitations.append(
                f"{duplicate_rows} repeated identifier rows merged per exact identifier; "
                "all distinct published values and source rows were preserved"
            )
        if mode == "stale_official_cache":
            limitations.append(
                "live official workbook refresh failed; using newest parseable official cache"
            )
        self.last_provenance = {
            "mode": mode,
            "retrieved_at": retrieved_at.isoformat(timespec="seconds"),
            "source_url": WORKBOOK_URL,
            "landing_url": LANDING_URL,
            "source_title": snapshot["source_title"],
            "sheet": snapshot["sheet"],
            "header_row": snapshot["header_row"],
            "headers": snapshot["headers"],
            "data_as_of": snapshot["data_as_of"],
            "status": "partial" if mode == "stale_official_cache" else "complete",
            "complete": mode != "stale_official_cache",
            "stale": mode == "stale_official_cache",
            "total_available": snapshot["source_rows"],
            "records": len(snapshot["rows"]),
            "missing_identifier_rows": missing,
            "generated_identifier_rows": int(
                snapshot.get("generated_identifier_rows") or 0
            ),
            "duplicate_source_id_rows": duplicate_rows,
            "duplicate_source_ids": snapshot["duplicate_source_ids"],
            "sha256": snapshot["sha256"],
            "classification": "agency forecast; never an active notice",
        }
        if limitations:
            self.last_provenance["limitations"] = limitations
            self.last_provenance["limitation"] = "; ".join(limitations)
        if private_error:
            self.last_provenance["private_error"] = private_error

    def _load_snapshot(self) -> tuple[dict[str, Any], datetime]:
        cache = _cache_dir() / (
            f"army_acquisition_forecast_{_today().isoformat()}.xlsx"
        )
        if cache.exists():
            try:
                snapshot = parse_workbook(cache.read_bytes())
                observed = _retrieved_at(cache)
                self._provenance(
                    snapshot, mode="official_daily_cache", retrieved_at=observed
                )
                return snapshot, observed
            except (OSError, ValueError):
                pass
        try:
            payload = get_bytes(
                WORKBOOK_URL,
                headers=HEADERS,
                timeout=45.0,
                retries=2,
            )
            snapshot = parse_workbook(payload)
            _atomic_bytes(cache, payload)
            observed = _retrieved_at(cache)
            self._provenance(
                snapshot, mode="live_official_workbook", retrieved_at=observed
            )
            return snapshot, observed
        except Exception as live_error:  # noqa: BLE001 - named stale fallback
            latest = _latest_valid_cache({cache})
            if latest is None:
                raise
            path, snapshot = latest
            observed = _retrieved_at(path)
            self._provenance(
                snapshot,
                mode="stale_official_cache",
                retrieved_at=observed,
                private_error=str(live_error),
            )
            return snapshot, observed

    def healthcheck(self) -> tuple[bool, str]:
        if not self.enabled:
            return True, (
                "explicitly disabled by LILA_ENABLE_ARMY_ACQUISITION_FORECAST"
            )
        try:
            snapshot, _observed = self._load_snapshot()
            return True, (
                f"official Army workbook parsed ({len(snapshot['rows'])} "
                "identified forecast records)"
            )
        except Exception as exc:  # noqa: BLE001
            return False, f"official Army forecast workbook unavailable: {exc}"

    def forecasts(
        self, query: Optional[SourceQuery] = None
    ) -> list[ForecastRecord]:
        if not self.enabled:
            self.last_provenance = {
                "mode": "not_run",
                "complete": False,
                "records": 0,
                "source_url": WORKBOOK_URL,
                "classification": "agency forecast; never an active notice",
            }
            return []
        snapshot, observed = self._load_snapshot()
        output: list[ForecastRecord] = []
        mapping_errors = 0
        for row in snapshot["rows"]:
            try:
                mapped = map_row(
                    row,
                    fiscal_year=snapshot["fiscal_year"],
                    data_as_of=snapshot["data_as_of"],
                    retrieved_at=observed,
                )
            except Exception:  # noqa: BLE001 - one malformed line never sinks the book
                mapping_errors += 1
                continue
            if mapped.source_id and mapped.title:
                output.append(mapped)
            else:
                mapping_errors += 1
        self.last_provenance["records"] = len(output)
        if mapping_errors:
            self.last_provenance["mapping_errors"] = mapping_errors
            self.last_provenance["complete"] = False
            limitation = (
                f"{mapping_errors} identified source rows could not be normalized"
            )
            limitations = self.last_provenance.setdefault("limitations", [])
            if limitation not in limitations:
                limitations.append(limitation)
            self.last_provenance["limitation"] = "; ".join(limitations)
        return output
