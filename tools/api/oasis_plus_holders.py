"""OASIS+ contract-holder roster adapter - who can already receive orders (no key).

OASIS+ is GSA's governmentwide services MAC family. For competitive and teaming
context the weekly contractor list is the authoritative roster of who holds
which domain in which family, and it carries signal no award dataset has: a
per-contract "Contract Status" with a Dormant flag (dormant holders cannot
receive new task orders), the contract's PRIMARY NAICS, the holder's UEI, and
named COCM/COPM contractor POCs with emails. Weekly award deltas surface new
entrants months before their first task order appears in USAspending.

The file is a date-stamped XLSX republished weekly, so the current link must be
re-resolved from the landing page on every pull (a stale direct URL 404s after
rotation).

Landing page: https://www.gsa.gov/buy-through-us/products-and-services/professional-services/buy-services/oasis-plus
List file (example): https://www.gsa.gov/system/files/OASIS%2B%20Contractor%20list%2008062026.xlsx

Workbook shape (verified 2026-08-08 against the 2026-08-06 drop):
  - "OASIS+Contract Information": preamble row "Last Updated 8/6/2026", then a
    header row, then one row per contract with status/UEI/PRIMARY NAICS/POCs.
  - Six family sheets (8a, Small Business, Woman Owned SB, Unrestricted,
    Service Disabled Veteran Owned, HUBZone): one row per contract x domain x
    NAICS with IDIQ/Domain/Vendor/URL/Contract #/SIN/NAICS columns. Quirk: the
    SDVOSB sheet's Domain column holds the sentinel "OASIS+VO"; the real domain
    is recovered from the SIN (digits 2-3 are the domain code, learned from the
    sheets where Domain is populated).
  - "Mapping" sheet (SIN to NAICS) is tolerated and ignored.
  - Woman Owned SB caveat: most WOSB contracts appear ONLY on the family sheet
    (no UEI/status/POC row); they are kept as thin holders and disclosed.
"""

from __future__ import annotations

import io
import re
import time
import urllib.request
from datetime import datetime, timezone
from typing import Any, Callable, Dict, List, Mapping, Optional, Tuple

import openpyxl

from tools.api._http import get_text
from tools.api.base import (DataSource, SourceKind, SourceQuery,
                            cap_disclosed, register_source)
from tools.api.provenance import ProvenanceEnvelope

LANDING_URL = ("https://www.gsa.gov/buy-through-us/products-and-services/"
               "professional-services/buy-services/oasis-plus")
#: the date-stamped file link as it appears on the landing page
_LIST_LINK_RE = re.compile(
    r'href="(https://www\.gsa\.gov/[^"]*contractor[^"]*list[^"]*\.xlsx)"',
    re.IGNORECASE,
)
_LAST_UPDATED_RE = re.compile(r"last\s+updated\s+(\d{1,2})/(\d{1,2})/(\d{4})",
                              re.IGNORECASE)
#: SDVOSB sheet's Domain column carries this instead of a real domain name
SDVOSB_DOMAIN_SENTINEL = "OASIS+VO"
#: info-sheet Project ID -> family label used on the family sheets' IDIQ column
PROJECT_FAMILY = {
    "OAS+8A": "OASIS+ 8(a)",
    "OAS+SB": "OASIS+ SB",
    "OAS+WO": "OASIS+ WOSB",
    "OAS+UR": "OASIS+ UR",
    "OAS+DV": "OASIS+ SDVOSB",
    "OAS+HZ": "OASIS+ HUBZone",
}
DEFAULT_CAP = 200
MAX_XLSX_BYTES = 30 * 1024 * 1024  # weekly file is ~5.5 MB; refuse a runaway
DOWNLOAD_TIMEOUT = 120.0
_UA = "Mozilla/5.0 (compatible; lila-source-adapter)"


def _download_bytes(url: str, timeout: float = DOWNLOAD_TIMEOUT,
                    retries: int = 2) -> bytes:
    """Binary GET. The shared _http helpers speak JSON/text only; this one
    XLSX leg uses the stdlib with the same bounded-retry policy."""
    last_exc: Optional[Exception] = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": _UA})
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read(MAX_XLSX_BYTES + 1)
        except Exception as exc:  # noqa: BLE001 - re-raised after retry budget
            last_exc = exc
            if attempt < retries - 1:
                time.sleep(0.5 * (2 ** attempt))
    assert last_exc is not None
    raise last_exc


def _s(cell: Any) -> Optional[str]:
    """Cell to stripped string, empty/None collapsed to None."""
    if cell is None:
        return None
    text = str(cell).strip()
    return text or None


def _resolve_list_url(landing_html: str) -> Optional[str]:
    match = _LIST_LINK_RE.search(landing_html or "")
    return match.group(1) if match else None


def _grids_from_xlsx(data: bytes) -> Dict[str, List[list]]:
    """Workbook bytes to {sheet_name: rows}; trailing empty cells trimmed,
    fully empty rows dropped."""
    if not data or not data.startswith(b"PK"):
        raise ValueError(
            "response is not an XLSX (zip magic missing); the landing page "
            "link may have rotated or returned an error page")
    if len(data) > MAX_XLSX_BYTES:
        raise ValueError(
            "file exceeds %d bytes; refusing runaway download" % MAX_XLSX_BYTES)
    workbook = openpyxl.load_workbook(
        io.BytesIO(data), read_only=True, data_only=True)
    grids: Dict[str, List[list]] = {}
    try:
        for name in workbook.sheetnames:
            rows: List[list] = []
            for row in workbook[name].iter_rows(values_only=True):
                cells = list(row)
                while cells and cells[-1] is None:
                    cells.pop()
                if cells:
                    rows.append(cells)
            grids[name] = rows
    finally:
        workbook.close()
    return grids


def _classify_sheet(rows: List[list]) -> Tuple[str, int]:
    """('info'|'family'|'other', header_row_index)."""
    for index, row in enumerate(rows[:6]):
        cells = {str(c).strip().lower() for c in row if c is not None}
        if "contract number" in cells and "uei" in cells:
            return "info", index
        if "contract #" in cells and "sin" in cells:
            return "family", index
    return "other", -1


def _header_map(row: list) -> Dict[str, int]:
    return {str(c).strip().lower(): i for i, c in enumerate(row)
            if c is not None and str(c).strip()}


def _cell(row: list, cols: Dict[str, int], key: str) -> Optional[str]:
    index = cols.get(key, -1)
    if index < 0 or index >= len(row):
        return None
    return _s(row[index])


def _holders_from_grids(
        grids: Mapping[str, List[list]]) -> Tuple[List[dict], Dict[str, Any]]:
    """Normalize the workbook grids into one holder per contract number.

    Tolerates the preamble "Last Updated" row, ragged row widths, the SDVOSB
    Domain sentinel, and family-sheet-only (thin) holders. Malformed rows are
    dropped and counted, never fatal.
    """
    meta: Dict[str, Any] = {"data_as_of": None, "dropped_rows": 0,
                            "unresolved_domain_rows": 0,
                            "info_rows": 0, "family_rows": 0}
    holders: Dict[str, dict] = {}
    order: List[str] = []
    family_raw: List[dict] = []
    domain_by_code: Dict[str, str] = {}

    for rows in grids.values():
        kind, header_index = _classify_sheet(rows)
        if kind == "other":
            continue
        for row in rows[:header_index]:
            for cell in row:
                found = _LAST_UPDATED_RE.search(str(cell or ""))
                if found:
                    month, day, year = found.groups()
                    meta["data_as_of"] = "%s-%02d-%02d" % (
                        year, int(month), int(day))
        cols = _header_map(rows[header_index])
        for row in rows[header_index + 1:]:
            if kind == "info":
                contract = _cell(row, cols, "contract number")
                vendor = _cell(row, cols, "vendor name")
                if not contract or not vendor:
                    meta["dropped_rows"] += 1
                    continue
                meta["info_rows"] += 1
                status = _cell(row, cols, "contract status")
                pocs = []
                for role, name_key, email_key in (
                        ("COCM", "cocm", "cocm email"),
                        ("COPM", "copm", "copm email")):
                    name = _cell(row, cols, name_key)
                    email = _cell(row, cols, email_key)
                    if name or email:
                        pocs.append({"role": role, "name": name,
                                     "email": email})
                if contract not in holders:
                    order.append(contract)
                holders[contract] = {
                    "contract_number": contract,
                    "vendor": vendor,
                    "uei": _cell(row, cols, "uei"),
                    "project_id": _cell(row, cols, "project id"),
                    "family": PROJECT_FAMILY.get(
                        _cell(row, cols, "project id") or ""),
                    "status": status,
                    "dormant": (status or "").strip().lower() == "dormant",
                    "primary_naics": _cell(row, cols, "primary naics"),
                    "city": _cell(row, cols, "vendor city"),
                    "zip": _cell(row, cols, "zip code"),
                    "website": None,
                    "domains": [],
                    "naics": [],
                    "pocs": pocs,
                    "info_joined": True,
                    "_blob_extra": [],
                }
            else:
                contract = _cell(row, cols, "contract #")
                vendor = _cell(row, cols, "vendor")
                if not contract or not vendor:
                    meta["dropped_rows"] += 1
                    continue
                meta["family_rows"] += 1
                sin = _cell(row, cols, "sin") or ""
                domain = _cell(row, cols, "domain")
                if (domain and domain != SDVOSB_DOMAIN_SENTINEL
                        and len(sin) >= 3):
                    domain_by_code.setdefault(sin[1:3], domain)
                family_raw.append({
                    "contract": contract, "vendor": vendor,
                    "idiq": _cell(row, cols, "idiq"), "domain": domain,
                    "sin": sin, "naics": _cell(row, cols, "naics"),
                    "title": _cell(row, cols, "naics title"),
                    "url": _cell(row, cols, "url"),
                })

    for raw in family_raw:
        contract = raw["contract"]
        holder = holders.get(contract)
        if holder is None:
            holder = {
                "contract_number": contract, "vendor": raw["vendor"],
                "uei": None, "project_id": None, "family": None,
                "status": None, "dormant": None, "primary_naics": None,
                "city": None, "zip": None, "website": None,
                "domains": [], "naics": [], "pocs": [],
                "info_joined": False, "_blob_extra": [],
            }
            holders[contract] = holder
            order.append(contract)
        if raw["idiq"] and not holder["family"]:
            holder["family"] = raw["idiq"]
        if raw["url"] and not holder["website"]:
            holder["website"] = raw["url"]
        domain = raw["domain"]
        if not domain or domain == SDVOSB_DOMAIN_SENTINEL:
            resolved = domain_by_code.get(raw["sin"][1:3]) if len(
                raw["sin"]) >= 3 else None
            if resolved is None:
                meta["unresolved_domain_rows"] += 1
            domain = resolved
        if domain and domain not in holder["domains"]:
            holder["domains"].append(domain)
        if raw["naics"] and raw["naics"] not in holder["naics"]:
            holder["naics"].append(raw["naics"])
        if raw["title"]:
            holder["_blob_extra"].append(raw["title"])

    finished: List[dict] = []
    for contract in order:
        holder = holders[contract]
        blob_parts = [
            holder["vendor"], holder["family"], holder["project_id"],
            holder["contract_number"], holder["uei"], holder["city"],
            holder["primary_naics"],
        ] + holder["domains"] + holder["naics"] + holder.pop("_blob_extra")
        holder["_blob"] = " ".join(p for p in blob_parts if p).lower()
        finished.append(holder)
    return finished, meta


def _matches(holder: dict, query: SourceQuery) -> bool:
    """Narrowing filters: keywords (any-of, matches vendor/domain/family/etc.)
    AND naics_codes (any-of, prefix match on primary or awarded codes)."""
    if query.keywords:
        blob = holder["_blob"]
        if not any(k.lower().strip('"') in blob for k in query.keywords if k):
            return False
    if query.naics_codes:
        own = [c for c in [holder["primary_naics"]] + holder["naics"] if c]
        wanted = [str(c).strip() for c in query.naics_codes if str(c).strip()]
        if not any(code.startswith(q) for q in wanted for code in own):
            return False
    return True


class OasisPlusHoldersSource(DataSource):
    name = "oasis_plus_holders"
    kind = SourceKind.ENRICHMENT

    def __init__(self,
                 fetch_text: Optional[Callable[..., str]] = None,
                 fetch_bytes: Optional[Callable[..., bytes]] = None) -> None:
        self._fetch_text = fetch_text or get_text
        self._fetch_bytes = fetch_bytes or _download_bytes
        # weekly file: memoize one parse per resolved (date-stamped) URL so a
        # multi-call press run downloads the 5.5 MB workbook once
        self._memo_url: Optional[str] = None
        self._memo: Optional[Tuple[List[dict], Dict[str, Any]]] = None

    def search(self, query: SourceQuery) -> list:
        raise NotImplementedError("enrichment-only source; use holders()")

    def healthcheck(self) -> Tuple[bool, str]:
        try:
            html = self._fetch_text(LANDING_URL, timeout=8.0, retries=1)
        except Exception as exc:  # noqa: BLE001 - honest light, never a crash
            return False, "landing page unreachable: %s" % exc
        list_url = _resolve_list_url(html)
        if not list_url:
            return False, ("landing page reachable but no contractor list "
                           ".xlsx link found (page layout may have changed)")
        return True, "contractor list link present (%s)" % list_url.rsplit(
            "/", 1)[-1]

    def enrich(self, query: SourceQuery) -> Dict[str, Any]:
        return self.holders(query)

    def holders(self, query: SourceQuery) -> Dict[str, Any]:
        """OASIS+ holders matching the query's keywords/NAICS lanes.

        Keywords match vendor names, domain names (e.g. "intelligence",
        "logistics"), family labels, NAICS titles, contract numbers, and UEIs.
        No keywords/NAICS means the full roster, capped and disclosed. The
        whole-market summary block is computed before any cap so a focused
        slice never reads as a dead market.
        """
        retrieved_at = datetime.now(timezone.utc)
        try:
            html = self._fetch_text(LANDING_URL, timeout=30.0, retries=2)
            list_url = _resolve_list_url(html)
            if not list_url:
                return self._failure(
                    "no contractor list .xlsx link on the OASIS+ landing "
                    "page; layout may have changed", retrieved_at)
            if self._memo_url != list_url or self._memo is None:
                data = self._fetch_bytes(list_url)
                self._memo = _holders_from_grids(_grids_from_xlsx(data))
                self._memo_url = list_url
            all_holders, meta = self._memo
        except Exception as exc:  # noqa: BLE001 - failure isolation contract
            return self._failure(str(exc), retrieved_at)

        matched = [h for h in all_holders if _matches(h, query)]
        cap = max(1, min(int(query.limit or DEFAULT_CAP), DEFAULT_CAP))
        kept = [{k: v for k, v in h.items() if k != "_blob"}
                for h in matched[:cap]]
        out = cap_disclosed(matched, cap, key="holders")
        out["holders"] = kept

        dormant = sum(1 for h in all_holders if h["dormant"])
        by_family: Dict[str, int] = {}
        for holder in all_holders:
            label = holder["family"] or "unstated"
            by_family[label] = by_family.get(label, 0) + 1
        thin = sum(1 for h in all_holders if not h["info_joined"])
        out["summary"] = {
            "total_holders": len(all_holders),
            "active": sum(1 for h in all_holders
                          if (h["status"] or "").lower() == "active"),
            "dormant": dormant,
            "by_family": by_family,
            "family_sheet_only": thin,
        }
        limitations = [
            "dormant holders cannot receive new task orders; the dormant "
            "flag is per contract in the source file",
        ]
        if thin:
            limitations.append(
                "%d holders appear only on a family sheet (mostly WOSB) and "
                "carry no UEI/status/POC columns in the source workbook"
                % thin)
        if meta["dropped_rows"]:
            limitations.append("%d malformed rows dropped during parse"
                               % meta["dropped_rows"])
        if meta["unresolved_domain_rows"]:
            limitations.append(
                "%d family rows had no resolvable domain (SIN code absent "
                "from the file's populated Domain columns)"
                % meta["unresolved_domain_rows"])
        out["list_url"] = list_url
        out["landing_url"] = LANDING_URL
        out["data_last_updated"] = meta["data_as_of"]
        out["_provenance"] = ProvenanceEnvelope(
            source=self.name,
            status="complete",
            mode="live_gsa_oasis_plus_xlsx",
            retrieval_mode="live",
            retrieved_at=retrieved_at,
            data_as_of=meta["data_as_of"],
            record_count=len(kept),
            limitations=limitations,
            public_detail="GSA OASIS+ weekly contractor list (last updated "
                          "%s)" % (meta["data_as_of"] or "unstated"),
        ).model_dump(mode="json")
        return out

    def _failure(self, message: str,
                 retrieved_at: datetime) -> Dict[str, Any]:
        return {
            "holders": [],
            "total_matched": 0,
            "error": message,
            "landing_url": LANDING_URL,
            "_provenance": ProvenanceEnvelope(
                source=self.name,
                status="failed",
                mode="live_gsa_oasis_plus_xlsx",
                retrieval_mode="live",
                retrieved_at=retrieved_at,
                record_count=0,
                limitations=[message],
            ).model_dump(mode="json"),
        }


# Registration: register_source requires a SourceSpec row in
# tools/api/source_catalog.py, which is the lead's wiring step (this build may
# not edit shared modules). The guard keeps the module importable today;
# registration takes effect the moment the catalog row lands.
try:
    register_source(OasisPlusHoldersSource)
except ValueError:
    pass
