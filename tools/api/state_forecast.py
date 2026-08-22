"""Department of State Procurement Forecast adapter (forecast-stage discovery).

State does not publish to the GSA Acquisition Gateway; its statutorily
required procurement forecast (P.L. 100-656) ships as one xlsx on state.gov.
The workbook carries fields no other pipeline source has for State before a
solicitation exists: 'Incumbent Contractor Name', 'Awarded Contract Order'
(real 19-prefix State PIIDs on roughly a third of rows, joinable to
FPDS/USAspending), 'Estimated Solicitation Date', 'Set-Aside Type',
'Acquisition Phase', and named POC / small-business-specialist emails.
That is recompete targeting with named incumbents months before anything
posts to SAM.gov.

Landing page: https://www.state.gov/procurement-forecast/
Workbook:     linked from the landing page as FY26-Procurement-Forecast-N.xlsx
              (verified 2026-08-08: FY26-Procurement-Forecast-4.xlsx). The
              filename embeds the revision number, so every pull re-discovers
              the current link from the landing page instead of pinning a URL.

Operational facts, verified live 2026-08-08:
- state.gov serves a block page to non-browser user agents; every request
  here sends a Chrome User-Agent through the _http helpers' headers param.
- Forecast rows are stated agency intent, never live deals: posted_date and
  response_deadline stay None, and the acquisition phase rides in
  raw_payload so downstream can see how early each line is.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import time
from datetime import date, datetime, timezone
from email.utils import parsedate_to_datetime
from typing import Any, Callable, Optional
from urllib.parse import urljoin

from openpyxl import load_workbook

from agents.schemas import OpportunityContact, RawOpportunity
from tools.api import _http
from tools.api.base import (
    DataSource,
    SourceKind,
    SourceQuery,
    cap_disclosed,
    register_source,
)
from tools.api.provenance import ProvenanceEnvelope
from tools.text_match import matching_phrases

LOG = logging.getLogger(__name__)

SOURCE_NAME = "state_forecast"
AGENCY = "Department of State"
LANDING_URL = "https://www.state.gov/procurement-forecast/"

#: state.gov fronts a WAF that returns a "Technical Difficulties" page to
#: non-browser user agents. A plain Chrome UA is sufficient; no cookies or
#: JS are needed (verified 2026-08-05 registry probe and again 2026-08-08).
BROWSER_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/126.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

_XLSX_HREF_RE = re.compile(
    r"""href=["']([^"']*procurement[-_ ]forecast[^"']*\.xlsx)["']""",
    re.IGNORECASE,
)
_NAICS_RE = re.compile(r"(?<!\d)\d{6}(?!\d)")
#: State Department PIIDs observed in the file: 19AQMM25C0285, 191NLE24C0002.
_PIID_RE = re.compile(r"^19[A-Z0-9]{6,}$")
_MONEY_RE = re.compile(r"\$\s*(\d[\d,]*(?:\.\d+)?)\s*([KMB])?", re.IGNORECASE)
_MULT = {"K": 1e3, "M": 1e6, "B": 1e9}

#: Placeholder spellings witnessed in the real FY26 workbook, including the
#: 'Dertermined' typo on live rows. Typed fields treat these as absent; the
#: verbatim value still rides in raw_payload for traceability.
_PLACEHOLDERS = {
    "", "tbd", "none", "n/a", "na", "not applicable",
    "to be determined", "to be dertermined",
}

#: Verbatim workbook headers (canonicalized) to raw_payload keys. Both the
#: 'Acquistion Phase' typo shipped in the real file and the corrected
#: spelling are mapped so a fixed future revision keeps parsing.
_HEADER_KEYS = {
    "requirement title": "requirement_title",
    "requirement description": "requirement_description",
    "agency": "agency",
    "contracting office": "contracting_office",
    "place of performance (city)": "pop_city",
    "place of performance (state)": "pop_state",
    "place of performance (country)": "pop_country",
    "naics code": "naics_raw",
    "acquistion phase": "acquisition_phase",
    "acquisition phase": "acquisition_phase",
    "estimated contract value*": "estimated_contract_value",
    "estimated contract value": "estimated_contract_value",
    "estimated award fy quarter": "estimated_award_fy_quarter",
    "estimated solicitation date": "estimated_solicitation_date",
    "pop date start": "pop_date_start",
    "extent competed": "extent_competed",
    "fiscal year": "fiscal_year",
    "set-aside type": "set_aside_type",
    "set aside type": "set_aside_type",
    "contract type": "contract_type",
    "incumbent contractor name": "incumbent_contractor_name",
    "awarded contract order": "awarded_contract_order",
    "anticipated award type": "anticipated_award_type",
    "facility security clearance": "facility_security_clearance",
    "poc name": "poc_name",
    "poc email": "poc_email",
    "small business specialist name": "small_business_specialist_name",
    "small business specialist email": "small_business_specialist_email",
}


# --------------------------------------------------------------------------- #
# Fetch layer
# --------------------------------------------------------------------------- #
def _get_bytes(
    url: str,
    *,
    headers: Optional[dict] = None,
    timeout: float = 60.0,
    retries: int = 3,
) -> "tuple[bytes, dict]":
    """GET raw bytes plus response headers, with _http's bounded-retry shape.

    _http exposes json/text helpers only, and an xlsx is a zip container
    that text decoding would corrupt. Rather than import httpx here, reuse
    _http's configured client factory (IPv4 pin, redirects, timeout) inside
    the same 0.5s/1s/2s backoff loop get_text uses.
    """
    last_exc: Optional[Exception] = None
    for attempt in range(retries):
        try:
            with _http._client(timeout) as client:
                resp = client.get(url, headers=headers)
            resp.raise_for_status()
            return resp.content, dict(resp.headers)
        except Exception as exc:  # noqa: BLE001 - transport error family varies
            last_exc = exc
            if attempt < retries - 1:
                time.sleep(0.5 * (2 ** attempt))
    assert last_exc is not None
    raise last_exc


def discover_workbook_url(html: str) -> Optional[str]:
    """First procurement-forecast .xlsx link on the landing page, absolute."""
    match = _XLSX_HREF_RE.search(html or "")
    if not match:
        return None
    return urljoin(LANDING_URL, match.group(1))


# --------------------------------------------------------------------------- #
# Parse layer
# --------------------------------------------------------------------------- #
def _canon_header(value: Any) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split()).lower()


def _fallback_key(canon: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", canon).strip("_")


def _clean_cell(value: Any) -> Any:
    """Collapse whitespace/nbsp in strings; empty strings become None."""
    if isinstance(value, str):
        text = " ".join(value.replace("\xa0", " ").split())
        return text or None
    return value


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    cleaned = _clean_cell(value)
    if cleaned is None:
        return None
    return cleaned if isinstance(cleaned, str) else str(cleaned)


def _stated(value: Any) -> Optional[str]:
    """A real stated value, or None when the cell is a TBD-style placeholder."""
    text = _text(value)
    if text is None or text.lower() in _PLACEHOLDERS:
        return None
    return text


def _rows_from_workbook(blob: bytes) -> "list[dict]":
    """Header-keyed row dicts from the first worksheet, dates as ISO strings.

    This is the exact shape the recorded JSON fixture stores, so the fixture
    slots into the normalize path where the xlsx parse leaves off.
    """
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        header_cells = next(rows_iter, None)
        if not header_cells:
            raise ValueError("forecast workbook has no header row")
        header = [str(cell) if cell is not None else "" for cell in header_cells]
        canon = {_canon_header(name) for name in header}
        if "requirement title" not in canon:
            raise ValueError(
                "unrecognized forecast workbook layout: "
                "no 'Requirement Title' column in the first sheet"
            )
        rows: "list[dict]" = []
        for cells in rows_iter:
            if cells is None or all(cell is None for cell in cells):
                continue
            row: dict = {}
            for index, name in enumerate(header):
                value = cells[index] if index < len(cells) else None
                if isinstance(value, datetime):
                    value = value.date().isoformat()
                elif isinstance(value, date):
                    value = value.isoformat()
                row[name] = value
            rows.append(row)
        return rows
    finally:
        workbook.close()


def _keyed(row: dict) -> dict:
    out: dict = {}
    for raw_name, value in row.items():
        canon = _canon_header(raw_name)
        key = _HEADER_KEYS.get(canon) or _fallback_key(canon)
        if key:
            out[key] = _clean_cell(value)
    return out


def _naics_codes(value: Any) -> "list[str]":
    """All 6-digit codes in the cell ('541512/541519' and int cells alike)."""
    stated = _stated(value)
    if not stated:
        return []
    return _NAICS_RE.findall(stated)


def _value_lower_bound(band: Any) -> Optional[float]:
    """Lower bound of the stated dollar band, in dollars.

    Matches the house rule (tools/api/forecasts/store.py value_rank): only
    numbers explicitly present in the source string become pipeline dollars.
    '$50.1M - $100M' -> 50_100_000.0. A lower token missing its magnitude
    suffix inherits the upper's ('$10.1 - $20M' appears in the real file).
    'Below $150K' states no floor, so it parses to None.
    """
    text = _stated(band)
    if not text:
        return None
    if text.lower().startswith(("below", "less than", "under")):
        return None
    tokens = _MONEY_RE.findall(text)
    if not tokens:
        return None
    number, scale = tokens[0]
    if not scale:
        for _, later_scale in tokens[1:]:
            if later_scale:
                scale = later_scale
                break
    if not scale:
        return None
    return float(number.replace(",", "")) * _MULT[scale.upper()]


def _piid(value: Any) -> Optional[str]:
    """The awarded contract order, only when it looks like a real State PIID."""
    stated = _stated(value)
    if not stated:
        return None
    candidate = stated.upper().replace(" ", "")
    return candidate if _PIID_RE.match(candidate) else None


def _stable_id(
    title: str,
    office: Optional[str],
    fiscal_year: Optional[str],
    quarter: Optional[str],
    ordinal: int,
) -> str:
    """Stable hash of (title + office + FY quarter), 'state-fcst-' prefixed.

    The workbook genuinely repeats lines (one FY26 title appears seven times
    in the same office and quarter), so repeated keys take a file-order
    ordinal instead of silently colliding in per-source dedupe.
    """
    parts = [
        title.lower(),
        (office or "").lower(),
        (fiscal_year or "").lower(),
        (quarter or "").lower(),
    ]
    if ordinal > 1:
        parts.append(str(ordinal))
    digest = hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()[:16]
    return "state-fcst-" + digest


# --------------------------------------------------------------------------- #
# Query translation + normalize
# --------------------------------------------------------------------------- #
def _matches(rec: dict, codes: "list[str]", query: SourceQuery) -> bool:
    if query.naics_codes:
        wanted = [code.strip() for code in query.naics_codes if code and code.strip()]
        if wanted and not any(
            code.startswith(prefix) for code in codes for prefix in wanted
        ):
            return False
    if query.agencies:
        hay = " ".join(
            filter(
                None,
                [AGENCY, _text(rec.get("agency")), _text(rec.get("contracting_office"))],
            )
        ).lower()
        if not any(agency.lower() in hay for agency in query.agencies if agency):
            return False
    if query.set_asides:
        set_aside = (_text(rec.get("set_aside_type")) or "").lower()
        if not any(term.lower() in set_aside for term in query.set_asides if term):
            return False
    if query.keywords:
        hay = " ".join(
            filter(
                None,
                [
                    _text(rec.get("requirement_title")),
                    _text(rec.get("requirement_description")),
                    _text(rec.get("contracting_office")),
                    _text(rec.get("incumbent_contractor_name")),
                    _text(rec.get("contract_type")),
                    _text(rec.get("anticipated_award_type")),
                ],
            )
        )
        if not matching_phrases(hay, query.keywords):
            return False
    return True


def _contacts(rec: dict) -> "list[OpportunityContact]":
    contacts: "list[OpportunityContact]" = []
    poc_name, poc_email = _stated(rec.get("poc_name")), _stated(rec.get("poc_email"))
    if poc_name or poc_email:
        contacts.append(
            OpportunityContact(name=poc_name, email=poc_email, contact_type="primary")
        )
    sbs_name = _stated(rec.get("small_business_specialist_name"))
    sbs_email = _stated(rec.get("small_business_specialist_email"))
    if sbs_name or sbs_email:
        contacts.append(
            OpportunityContact(
                name=sbs_name,
                email=sbs_email,
                contact_type="small_business_specialist",
            )
        )
    return contacts


def _normalize(
    rows: "list[dict]", workbook_url: str, query: SourceQuery
) -> "tuple[list[RawOpportunity], int]":
    """(matched opportunities in file order, dropped malformed-row count).

    Stable ids are assigned over every title-bearing row BEFORE filtering,
    so a narrower query never shifts which duplicate line owns which id.
    """
    matched: "list[RawOpportunity]" = []
    dropped = 0
    seen: "dict[tuple, int]" = {}
    for raw in rows:
        rec = _keyed(raw)
        title = _stated(rec.get("requirement_title"))
        if not title:
            dropped += 1
            continue
        office = _stated(rec.get("contracting_office"))
        fiscal_year = _stated(rec.get("fiscal_year"))
        quarter = _stated(rec.get("estimated_award_fy_quarter"))
        key = (
            title.lower(),
            (office or "").lower(),
            (fiscal_year or "").lower(),
            (quarter or "").lower(),
        )
        seen[key] = seen.get(key, 0) + 1
        codes = _naics_codes(rec.get("naics_raw"))
        if not _matches(rec, codes, query):
            continue
        payload = dict(rec)
        payload.update(
            {
                "naics_codes_all": codes,
                "piid": _piid(rec.get("awarded_contract_order")),
                "source_url": workbook_url,
                "landing_url": LANDING_URL,
            }
        )
        matched.append(
            RawOpportunity(
                source=SOURCE_NAME,
                source_id=_stable_id(title, office, fiscal_year, quarter, seen[key]),
                title=title,
                agency=AGENCY,
                naics_code=codes[0] if codes else None,
                set_aside=_stated(rec.get("set_aside_type")),
                posted_date=None,
                response_deadline=None,
                estimated_value=_value_lower_bound(rec.get("estimated_contract_value")),
                api_url=workbook_url,
                contacts=_contacts(rec),
                raw_payload=payload,
            )
        )
    return matched, dropped


def _data_as_of(headers: Optional[dict]) -> Optional[str]:
    """The workbook's own Last-Modified date; never synthesized from now()."""
    stamp = None
    for name, value in (headers or {}).items():
        if str(name).lower() == "last-modified":
            stamp = value
            break
    if not stamp:
        return None
    try:
        return parsedate_to_datetime(str(stamp)).date().isoformat()
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------- #
# Adapter
# --------------------------------------------------------------------------- #
class StateForecastSource(DataSource):
    name = SOURCE_NAME
    # Planning intent is never eligible for the canonical live-opportunity
    # discover() fan-out.  RawOpportunity is only an internal normalization
    # surface bridged immediately to ForecastRecord by forecasts().
    kind = SourceKind.ENRICHMENT
    enabled = True

    def __init__(
        self,
        fetch_text: Optional[Callable[..., str]] = None,
        fetch_bytes: Optional[Callable[..., "tuple[bytes, dict]"]] = None,
    ) -> None:
        from tools.toggles import is_enabled
        self.enabled = is_enabled(self.name, True)
        self.last_provenance: "Optional[dict]" = None
        # Injection seams for tests. Defaults resolve at call time through
        # the _http module attribute so the suite's live-HTTP receipt seam
        # observes real calls.
        self._fetch_text = fetch_text
        self._fetch_bytes = fetch_bytes

    def healthcheck(self) -> "tuple[bool, str]":
        fetch_text = self._fetch_text or _http.get_text
        try:
            html = fetch_text(
                LANDING_URL, headers=BROWSER_HEADERS, timeout=8.0, retries=1
            )
        except Exception as exc:  # noqa: BLE001 - probe reports, never raises
            return False, f"state.gov landing page unreachable: {exc}"
        workbook_url = discover_workbook_url(html)
        if not workbook_url:
            return False, (
                "landing page loads but no procurement-forecast .xlsx link "
                "found (layout drift or WAF block page)"
            )
        return True, (
            "landing page live; current workbook: "
            + workbook_url.rsplit("/", 1)[-1]
        )

    def search(self, query: SourceQuery) -> "list[RawOpportunity]":
        fetch_text = self._fetch_text or _http.get_text
        fetch_bytes = self._fetch_bytes or _get_bytes
        try:
            html = fetch_text(LANDING_URL, headers=BROWSER_HEADERS)
            workbook_url = discover_workbook_url(html)
            if not workbook_url:
                raise ValueError(
                    "landing page has no procurement-forecast .xlsx link"
                )
            blob, resp_headers = fetch_bytes(workbook_url, headers=BROWSER_HEADERS)
            blob = bytes(blob or b"")
            if not blob.startswith(b"PK"):
                raise ValueError(
                    "workbook fetch returned non-xlsx bytes "
                    "(state.gov WAF block page?)"
                )
            rows = _rows_from_workbook(blob)
        except Exception as exc:  # noqa: BLE001 - failure isolation contract
            LOG.warning("state_forecast search failed: %s", exc)
            self.last_provenance = ProvenanceEnvelope(
                source=SOURCE_NAME,
                status="failed",
                mode="live_state_forecast_xlsx",
                retrieval_mode="live",
                retrieved_at=datetime.now(timezone.utc),
                record_count=0,
                limitations=[f"forecast retrieval failed: {exc}"],
            ).model_dump(mode="json")
            return []

        matched, dropped = _normalize(rows, workbook_url, query)
        limit = query.limit if query.limit and query.limit > 0 else len(matched)
        capped = cap_disclosed(matched, limit, key="items")
        kept: "list[RawOpportunity]" = capped["items"]

        limitations = [
            "Forecast lines are stated agency intent (P.L. 100-656 planning "
            "data), not live solicitations; entries are revised or cancelled "
            "routinely",
        ]
        if dropped:
            limitations.append(
                f"dropped {dropped} rows with no requirement title"
            )
        if capped.get("truncated"):
            limitations.append(str(capped.get("truncated_note")))
        envelope = ProvenanceEnvelope(
            source=SOURCE_NAME,
            status="partial" if dropped else "complete",
            mode="live_state_forecast_xlsx",
            retrieval_mode="live",
            retrieved_at=datetime.now(timezone.utc),
            data_as_of=_data_as_of(resp_headers),
            record_count=len(kept),
            limitations=limitations,
            public_detail=(
                "Live Department of State procurement forecast workbook ("
                + workbook_url.rsplit("/", 1)[-1]
                + ")"
            ),
            retrieval_url=workbook_url,
            retrieval_query=query.model_dump(mode="json"),
            raw_content_sha256=hashlib.sha256(blob).hexdigest(),
            raw_content_kind="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            normalization_version="state_forecast.v1.2026-08-08",
        )
        self.last_provenance = envelope.model_dump(mode="json")
        self.last_provenance.update({
            "source_url": workbook_url,
            "landing_url": LANDING_URL,
        })
        for opportunity in kept:
            opportunity.raw_payload["_provenance"] = self.last_provenance
        return kept

    def forecasts(self):
        from tools.api.forecasts.bridge import from_raw
        rows = self.search(SourceQuery(limit=5000))
        return [from_raw(row, self.last_provenance) for row in rows]


try:
    register_source(StateForecastSource)
except ValueError:
    # "state_forecast" has no SourceSpec row in tools/api/source_catalog.py
    # yet. The integration lead adds the catalog row and the registrar import
    # together; guarding here keeps the module importable and testable now
    # without this lane editing shared files.
    pass
