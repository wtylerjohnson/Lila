"""NASA SEWP VI awardee-roster adapter.

The roster establishes that a named prime appears in an awarded category.  It
does not establish a UEI, incumbency, order, or customer relationship.
"""

from __future__ import annotations

import io
from typing import Any, Callable, Mapping, Optional

from tools.api import _http
from tools.api._connector_wave3 import (failure, matches, result, scoped_terms,
                                        stable_id, text)
from tools.api.base import DataSource, SourceKind, SourceQuery, register_source

SOURCE_NAME = "nasa_sewp_vi"
AWARDEES_URL = "https://www.sewp.nasa.gov/documents/SEWPVI_Awardees.xlsx"
MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_RETURNED = 500


def _grids(blob: bytes) -> dict[str, list[list[Any]]]:
    if not blob.startswith(b"PK") or len(blob) > MAX_XLSX_BYTES:
        raise ValueError("SEWP VI response is not a bounded XLSX workbook")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - production dependency
        raise RuntimeError("openpyxl is required to read the official roster") from exc
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        return {
            sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()


def _normalize(grids: Mapping[str, list[list[Any]]]) -> list[dict[str, Any]]:
    by_name: dict[str, dict[str, Any]] = {}
    for sheet_name, rows in grids.items():
        category = text(sheet_name)
        header = -1
        name_col = -1
        for index, row in enumerate(rows[:10]):
            for col, value in enumerate(row):
                if (text(value) or "").casefold() in {"prime offeror name", "prime offeror"}:
                    header, name_col = index, col
                    break
            if header >= 0:
                break
        if header < 0:
            continue
        for row in rows[header + 1:]:
            if name_col >= len(row):
                continue
            name = text(row[name_col])
            if not name:
                continue
            key = name.casefold()
            record = by_name.setdefault(key, {
                "record_id": stable_id(SOURCE_NAME, name),
                "prime_offeror_name": name,
                "categories": [],
                "uei": None,
                "piids": [],
                "_join_keys": {"legal_name_inference": name},
                "evidence_semantics": "vehicle_access_not_incumbency",
            })
            if category and category not in record["categories"]:
                record["categories"].append(category)
    if not by_name:
        raise ValueError("no Prime Offeror Name columns found in SEWP VI workbook")
    return list(by_name.values())


class NasaSewpViSource(DataSource):
    name = SOURCE_NAME
    kind = SourceKind.ENRICHMENT

    def __init__(self, fetch_bytes: Optional[Callable[..., bytes]] = None,
                 parse_workbook: Optional[Callable[[bytes], Mapping[str, list[list[Any]]]]] = None) -> None:
        self._fetch_bytes = fetch_bytes or self._live_bytes
        self._parse_workbook = parse_workbook or _grids

    @staticmethod
    def _live_bytes(url: str, **kwargs: Any) -> bytes:
        with _http._client(float(kwargs.get("timeout", 30.0))) as client:
            response = client.get(url)
        response.raise_for_status()
        blob = bytes(response.content)
        if len(blob) > MAX_XLSX_BYTES:
            raise ValueError("SEWP VI workbook exceeded download bound")
        return blob

    def search(self, query: SourceQuery) -> list:
        raise NotImplementedError("enrichment-only source; use awardees()")

    def healthcheck(self) -> tuple[bool, str]:
        try:
            rows = _normalize(self._parse_workbook(
                self._fetch_bytes(AWARDEES_URL, timeout=15.0)))
            return True, f"official SEWP VI awardee workbook reachable ({len(rows)} primes)"
        except Exception as exc:  # noqa: BLE001
            return False, f"SEWP VI awardee workbook unavailable: {exc}"

    def awardees(self, query: SourceQuery) -> dict[str, Any]:
        terms = scoped_terms(query)
        if not terms:
            return result(
                self.name, items=[], status="partial", mode="refused_unscoped",
                url=AWARDEES_URL,
                limitations=["unscoped awardee-roster pull refused; pass a prime name or category"],
            )
        try:
            blob = self._fetch_bytes(AWARDEES_URL, timeout=30.0)
            native = _normalize(self._parse_workbook(blob))
        except Exception as exc:  # noqa: BLE001
            return failure(self.name, mode="live_official_xlsx", url=AWARDEES_URL,
                           error=exc)
        matched = [row for row in native if matches(row, terms)]
        cap = max(1, min(query.limit, MAX_RETURNED))
        kept = matched[:cap]
        limits = [
            "awardee name and category prove SEWP VI vehicle access only",
            "the workbook has no UEI or contract PIID; legal-name joins are inference-grade until independently resolved",
            "award roster may change during protest or corrective-action activity",
        ]
        if len(matched) > cap:
            limits.append(f"kept first {cap} of {len(matched)} scoped matches")
        return result(
            self.name, items=kept, total_matched=len(matched),
            status="partial" if len(matched) > cap else "complete",
            mode="live_official_xlsx_scoped_local_filter", limitations=limits,
            url=AWARDEES_URL, query={"terms": terms, "limit": cap}, raw=blob,
            content_kind="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


try:
    register_source(NasaSewpViSource)
except ValueError:
    pass
