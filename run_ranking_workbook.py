"""Discernment ranking workbook: the operator and the client rank fit.

Deterministic and offline: reads the pressed Pre-Assessment document and the
approved Analyst packet, writes one .xlsx with three ranking sheets
(opportunities, keywords, NAICS) carrying dropdown fit scales and notes
columns. Zero LLM, zero network; generated from pipeline truth, never
hand-authored. The recorded rankings are calibration input for the
relevance engine, entered by humans in the session.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Optional, Sequence

_ROOT = Path(__file__).resolve().parent
_STAGE = "ranking-workbook"

_NAVY = "12233A"
_RULE = "D9E1E8"


def _client_id(client: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "-", client.casefold()).strip("-")
    if not value:
        raise ValueError("client identity cannot produce a workbook id")
    return value


def _slug_underscore(client: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "_", client.casefold()).strip("_")
    return value


def build_ranking_workbook(
    client: str,
    *,
    root: Path | str = _ROOT,
    out: Path | str | None = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from agents.candidate_review_v1.contracts import CandidateReviewDocument

    root_path = Path(root)
    client_id = _client_id(client)
    slug = _slug_underscore(client)

    document_path = next(
        (path for candidate in (client_id, slug)
         for path in [root_path / "data" / "state" / "candidate_review_v1"
                      / candidate
                      / f"{candidate}.candidate_review.document.json"]
         if path.is_file()),
        None)
    if document_path is None:
        raise ValueError(
            "no pressed Pre-Assessment document; run the candidate_review "
            "step first")
    document = CandidateReviewDocument.model_validate_json(
        document_path.read_text(encoding="utf-8"))

    packet_path = next(
        (path for candidate in (slug, client_id)
         for path in [root_path / "data" / "review"
                      / f"{candidate}.review.json"]
         if path.is_file()),
        None)
    packet = (json.loads(packet_path.read_text(encoding="utf-8"))
              if packet_path is not None else {})
    strategy = packet.get("strategy") or {}

    evidence_by_id = {item.evidence_id: item for item in document.evidence}

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=_NAVY)
    wrap = Alignment(wrap_text=True, vertical="top")

    def _sheet(title, headers, widths):
        sheet = (workbook.active if workbook.sheetnames == ["Sheet"]
                 else workbook.create_sheet())
        sheet.title = title
        sheet.sheet_properties.tabColor = _NAVY
        for index, (header, width) in enumerate(zip(headers, widths), start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.font = header_font
            cell.fill = header_fill
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        return sheet

    def _fit_validation(sheet, column, rows, formula):
        if rows <= 0:
            return
        validation = DataValidation(
            type="list", formula1=formula, allow_blank=True,
            showDropDown=False)
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}{rows + 1}")

    # -- Sheet 1: opportunities -------------------------------------------
    opportunities = _sheet(
        "Opportunities",
        ("Opportunity", "Agency", "Kind", "Next date", "Source",
         "Ed fit (1-5)", "GTM fit (1-5)", "Notes"),
        (52, 30, 20, 13, 14, 12, 12, 44))
    row = 1
    for candidate in document.candidates:
        row += 1
        next_date = min(
            (item.sort_date for item in candidate.dates), default=None)
        urls = []
        for member in candidate.members:
            for evidence_id in member.member_evidence_ids:
                record = evidence_by_id.get(evidence_id)
                if record is not None:
                    urls.append(str(record.source_url))
        opportunities.cell(row=row, column=1,
                           value=candidate.title).alignment = wrap
        opportunities.cell(row=row, column=2, value=candidate.agency)
        opportunities.cell(
            row=row, column=3,
            value=candidate.kind.value.replace("_", " "))
        opportunities.cell(
            row=row, column=4,
            value=next_date.isoformat() if next_date else "")
        if urls:
            link = opportunities.cell(row=row, column=5, value="open record")
            link.hyperlink = urls[0]
            link.font = Font(color="157EAF", underline="single")
    _fit_validation(opportunities, "F", row - 1, '"1,2,3,4,5"')
    _fit_validation(opportunities, "G", row - 1, '"1,2,3,4,5"')

    # -- Sheet 2: keywords -------------------------------------------------
    keywords = _sheet(
        "Keywords",
        ("Keyword", "Category", "Rationale", "Verdict", "Notes"),
        (32, 18, 56, 16, 40))
    row = 1
    for keyword in strategy.get("keywords") or []:
        row += 1
        keywords.cell(row=row, column=1, value=keyword.get("term"))
        keywords.cell(row=row, column=2, value=keyword.get("category"))
        keywords.cell(row=row, column=3,
                      value=keyword.get("rationale")).alignment = wrap
    _fit_validation(keywords, "D", row - 1, '"Keep,Adjust,Drop"')

    # -- Sheet 3: NAICS ----------------------------------------------------
    naics = _sheet(
        "NAICS",
        ("Code", "Title", "Role", "Rationale", "Verdict", "Notes"),
        (10, 40, 12, 52, 16, 40))
    meta_by_code = {
        entry.get("code"): entry
        for entry in packet.get("naics_meta") or []}
    row = 1
    for code in strategy.get("inferred_naics") or []:
        row += 1
        entry = meta_by_code.get(code) or {}
        naics.cell(row=row, column=1, value=code)
        naics.cell(row=row, column=2, value=entry.get("title") or "")
        naics.cell(row=row, column=3, value=entry.get("role") or "")
        naics.cell(row=row, column=4,
                   value=entry.get("rationale") or "").alignment = wrap
    _fit_validation(naics, "E", row - 1, '"Keep,Adjust,Drop"')

    out_path = (Path(out) if out is not None
                else root_path / "data" / "reports"
                / f"{slug}.ranking_workbook.xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return out_path.resolve()


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate the discernment ranking workbook")
    parser.add_argument("--client", required=True)
    parser.add_argument("--out", default=None)
    args = parser.parse_args(argv)
    try:
        path = build_ranking_workbook(args.client, out=args.out)
    except Exception as exc:  # noqa: BLE001 - one named failure object
        sys.stderr.write(json.dumps({
            "stage": _STAGE,
            "reason": f"{type(exc).__name__}: {str(exc)[:300]}",
            "fix_surface": (
                f"run_candidate_review.py --client {args.client} --with-watch "
                "(the workbook ranks the pressed document)"),
        }, sort_keys=True) + "\n")
        return 2
    sys.stdout.write(f"[out:{_STAGE}] {path}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
