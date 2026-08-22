"""Department of State Procurement Forecast adapter: fixture-driven.

Fixture provenance (recorded, not invented):
    tests/fixtures/sources/state_forecast.json was fetched
    2026-08-08T13:19:37Z. Link discovery ran against the live landing page
    https://www.state.gov/procurement-forecast/ (Chrome UA; state.gov blocks
    non-browser user agents), which linked
    https://www.state.gov/wp-content/uploads/2026/04/FY26-Procurement-Forecast-4.xlsx
    (Last-Modified Thu, 16 Apr 2026 17:38:13 GMT). The workbook's first sheet
    'Acquisition Submission' held 453 data rows; the fixture keeps the REAL
    header plus the first 40 data rows in file order, datetime cells
    serialized to ISO dates, along with the REAL landing-page anchor block
    for link-discovery tests. The source's own typos ride along untouched
    ('Acquistion Phase' header, 'To Be Dertermined' cells, 'GDIT ' with a
    trailing space).

No network anywhere below except the single opt-in live smoke at the bottom.
The deterministic tests rebuild a genuine workbook in memory from the fixture
rows and inject it through the adapter's fetch seams, so the full
landing page -> link discovery -> bytes -> rows -> RawOpportunity path runs
exactly as in production.
"""

from __future__ import annotations

import io
import json
import logging
import os
from pathlib import Path

import openpyxl
import pytest

from tools.api.base import REGISTRY, SourceKind, SourceQuery
from tools.api.state_forecast import (
    LANDING_URL,
    StateForecastSource,
    _piid,
    _value_lower_bound,
    discover_workbook_url,
)

FIXTURE_PATH = (Path(__file__).parent / "fixtures" / "sources"
                / "state_forecast.json")
FIXTURE = json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))
HEADER = FIXTURE["header"]
ROWS = FIXTURE["rows"]
LANDING_HTML = FIXTURE["landing_html_excerpt"]

WORKBOOK_URL = ("https://www.state.gov/wp-content/uploads/2026/04/"
                "FY26-Procurement-Forecast-4.xlsx")
LAST_MODIFIED_HTTP = "Thu, 16 Apr 2026 17:38:13 GMT"
SOE_BPA_TITLE = ("Recompete: Blanket Purchase Agreement (BPA) for Global "
                 "Logistics - Specialized Operational Equip (SOE)")


def _xlsx_bytes(rows):
    """Rebuild a genuine workbook from fixture rows (same sheet name)."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Acquisition Submission"
    sheet.append(HEADER)
    for row in rows:
        sheet.append([row.get(name) for name in HEADER])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _source(rows=None, data=None, headers=None, html=LANDING_HTML):
    """Adapter with both fetch seams injected; no call leaves the machine."""
    payload = data if data is not None else _xlsx_bytes(
        ROWS if rows is None else rows)
    resp_headers = ({"Last-Modified": LAST_MODIFIED_HTTP}
                    if headers is None else headers)

    def fetch_text(url, **kw):
        return html

    def fetch_bytes(url, **kw):
        return payload, resp_headers

    return StateForecastSource(fetch_text=fetch_text, fetch_bytes=fetch_bytes)


# ------------------------------------------------------------------ #
# fixture integrity
# ------------------------------------------------------------------ #

def test_fixture_is_recorded_first_rows():
    prov = FIXTURE["_fixture_provenance"]
    assert prov["landing_url"] == LANDING_URL
    assert prov["workbook_url"] == WORKBOOK_URL
    assert prov["retrieved_at"] == "2026-08-08T13:19:37Z"
    assert prov["workbook_last_modified"] == LAST_MODIFIED_HTTP
    assert prov["total_data_rows_in_file"] == 453
    assert len(ROWS) == 40
    assert len(HEADER) == 25
    # every recorded row is full width, and the source's own header typo
    # shipped in the real file is preserved, never corrected
    assert all(set(row) == set(HEADER) for row in ROWS)
    assert "Acquistion Phase" in HEADER


# ------------------------------------------------------------------ #
# link discovery (the filename revs; every pull re-reads the landing page)
# ------------------------------------------------------------------ #

def test_discovers_workbook_url_from_recorded_anchor_block():
    assert discover_workbook_url(LANDING_HTML) == WORKBOOK_URL


def test_discovery_resolves_relative_links_and_handles_absence():
    relative = '<a href="/wp-content/uploads/FY27-Procurement-Forecast-1.xlsx">f</a>'
    assert discover_workbook_url(relative) == (
        "https://www.state.gov/wp-content/uploads/"
        "FY27-Procurement-Forecast-1.xlsx")
    assert discover_workbook_url("<html>Technical Difficulties</html>") is None
    assert discover_workbook_url("") is None


# ------------------------------------------------------------------ #
# field mapping
# ------------------------------------------------------------------ #

def test_maps_forecast_row_to_raw_opportunity():
    records = _source().search(SourceQuery())
    assert len(records) == 40          # arrival order, nothing invented
    first = records[0]
    assert first.source == "state_forecast"
    assert first.source_id.startswith("state-fcst-")
    assert len(first.source_id) == len("state-fcst-") + 16
    assert first.title == SOE_BPA_TITLE
    assert first.agency == "Department of State"
    assert first.naics_code == "334519"
    assert first.set_aside == "Small business-Total"
    # '$50.1M - $100M': only the stated lower bound becomes pipeline dollars
    assert first.estimated_value == 50_100_000.0
    assert str(first.api_url) == WORKBOOK_URL
    # forecast rows are stated intent, never live deals
    assert first.posted_date is None
    assert first.response_deadline is None
    payload = first.raw_payload
    assert payload["incumbent_contractor_name"] == "Culmen International, LLC"
    assert payload["awarded_contract_order"] == "To Be Determined"
    assert payload["estimated_solicitation_date"] == "1st Qtr FY26"
    assert payload["set_aside_type"] == "Small business-Total"
    assert payload["acquisition_phase"] == "Acquisition Planning"
    assert payload["fiscal_year"] == "FY26"
    assert payload["estimated_award_fy_quarter"] == "4th (July 1 - September 30)"
    assert payload["source_url"] == WORKBOOK_URL
    assert payload["landing_url"] == LANDING_URL
    # named POCs feed the contact plan
    contacts = {(c.contact_type, c.name, c.email) for c in first.contacts}
    assert ("primary", "Tony Sixon", "SixonAC@state.gov") in contacts
    assert ("small_business_specialist", "Tyler Sinclair",
            "SmallBusiness@state.gov") in contacts


def test_real_piid_rows_join_and_placeholders_stay_verbatim_in_payload():
    records = _source().search(SourceQuery())
    # 'Awarded Contract Order' carries a real 19-prefix State PIID on this
    # row: the FPDS/USAspending join key the registry verified in-file.
    gdit = next(r for r in records
                if r.raw_payload.get("piid") == "19AQMM25C0285")
    assert gdit.title == "ACN/CTR Global Programs"
    # the live cell reads 'GDIT ' with a trailing space; cleaned, not edited
    assert gdit.raw_payload["incumbent_contractor_name"] == "GDIT"
    # placeholder rows: typed fields go absent, the verbatim source text
    # (including the file's own 'Dertermined' typo) rides in raw_payload
    price = next(r for r in records if "(PRICE)" in r.title)
    assert price.set_aside is None
    assert price.raw_payload["piid"] is None
    assert price.raw_payload["incumbent_contractor_name"] == "To Be Dertermined"
    assert price.raw_payload["awarded_contract_order"] == "To Be Dertermined"


def test_value_lower_bound_follows_the_house_rule():
    # only numbers explicitly present in the source string become dollars
    assert _value_lower_bound("$50.1M - $100M") == 50_100_000.0
    assert _value_lower_bound("$251K - $500K") == 251_000.0
    # a lower token missing its magnitude suffix inherits the upper's
    assert _value_lower_bound("$10.1 - $20M") == 10_100_000.0
    # 'Below $150K' states no floor
    assert _value_lower_bound("Below $150K") is None
    assert _value_lower_bound("To Be Determined") is None
    assert _value_lower_bound(None) is None


def test_piid_shape_guard():
    assert _piid("19AQMM22C0072") == "19AQMM22C0072"
    assert _piid(" 19aqmm25c0285 ") == "19AQMM25C0285"
    assert _piid("To Be Determined") is None
    assert _piid("N/A") is None
    assert _piid("SAQMMA15F1234") is None      # not a 19-prefix State PIID


# ------------------------------------------------------------------ #
# stable ids
# ------------------------------------------------------------------ #

def test_source_id_stable_and_assigned_before_filtering():
    full = _source().search(SourceQuery())
    narrowed = _source().search(SourceQuery(naics_codes=["5415"]))
    # the live cell reads 'New Award:  Panama...' with a doubled space;
    # whitespace collapses, words never change
    assert [r.title for r in narrowed] == [
        "New Award: Panama Unified Information Sharing Targeting System - "
        "Border Information Analysis & Targeting Unit (BIATU) Project"]
    wide_match = next(r for r in full if r.naics_code == "541512")
    # a narrower query never shifts which row owns which id
    assert narrowed[0].source_id == wide_match.source_id
    assert len({r.source_id for r in full}) == len(full)


def test_repeated_lines_get_ordinal_ids_not_silent_collisions():
    records = _source(rows=[ROWS[0], ROWS[0]]).search(SourceQuery())
    assert len(records) == 2
    assert records[0].source_id != records[1].source_id
    # the first occurrence keeps the id it has when the line appears once
    solo = _source(rows=[ROWS[0]]).search(SourceQuery())
    assert records[0].source_id == solo[0].source_id


# ------------------------------------------------------------------ #
# provenance envelope
# ------------------------------------------------------------------ #

def test_every_payload_carries_provenance_envelope():
    records = _source().search(SourceQuery())
    assert records
    for opp in records:
        env = opp.raw_payload.get("_provenance")
        assert env, "payload without _provenance envelope"
        assert env["schema_version"] == 1
        assert env["source"] == "state_forecast"
        assert env["status"] == "complete"
        assert env["retrieval_mode"] == "live"
        assert env["record_count"] == len(records)
        assert env["retrieved_at"]
        # the workbook's own Last-Modified date, never retrieval time
        assert env["data_as_of"] == "2026-04-16"
        assert "P.L. 100-656" in " ".join(env["limitations"])
        assert "FY26-Procurement-Forecast-4.xlsx" in env["public_detail"]


def test_data_as_of_header_lookup_is_case_insensitive_and_optional():
    lower = _source(headers={"last-modified": LAST_MODIFIED_HTTP})
    records = lower.search(SourceQuery(limit=1))
    assert records[0].raw_payload["_provenance"]["data_as_of"] == "2026-04-16"
    bare = _source(headers={})
    records = bare.search(SourceQuery(limit=1))
    assert records[0].raw_payload["_provenance"]["data_as_of"] is None


# ------------------------------------------------------------------ #
# query filtering
# ------------------------------------------------------------------ #

def test_naics_prefix_filter():
    records = _source().search(SourceQuery(naics_codes=["5415"]))
    assert len(records) == 1
    assert records[0].naics_code == "541512"


def test_keyword_filter_reads_title_and_description():
    records = _source().search(SourceQuery(keywords=["logistics"]))
    assert [r.title for r in records] == [SOE_BPA_TITLE]


def test_set_aside_filter():
    records = _source().search(SourceQuery(set_asides=["8(a)"]))
    assert len(records) == 1
    assert records[0].set_aside == "8(a)"
    assert records[0].title == "FASTC Technical Staffing"


def test_agency_filter_matches_department_and_office_lines():
    assert len(_source().search(SourceQuery(agencies=["State"]))) == 40
    office = _source().search(SourceQuery(agencies=["ACN/CTR"]))
    assert len(office) == 11
    assert all(r.raw_payload["contracting_office"] == "ACN/CTR"
               for r in office)
    assert _source().search(SourceQuery(agencies=["DHS"])) == []


def test_no_match_returns_empty_list():
    assert _source().search(SourceQuery(keywords=["quantum-zzz-nomatch"])) == []


def test_limit_is_disclosed_not_silent():
    records = _source().search(SourceQuery(limit=5))
    assert len(records) == 5
    env = records[0].raw_payload["_provenance"]
    assert env["record_count"] == 5
    assert "kept first 5 of 40" in " ".join(env["limitations"])


# ------------------------------------------------------------------ #
# malformed tolerance and failure isolation
# ------------------------------------------------------------------ #

def test_malformed_rows_drop_and_are_counted():
    untitled = {name: None for name in HEADER}
    untitled["Contracting Office"] = "AQM"     # real cells, no title
    placeholder = dict(ROWS[1])
    placeholder["Requirement Title"] = "TBD"   # placeholder title = no title
    records = _source(rows=[untitled, ROWS[2], placeholder]).search(
        SourceQuery())
    assert len(records) == 1
    assert records[0].naics_code == "541512"      # the surviving BIATU row
    env = records[0].raw_payload["_provenance"]
    assert env["status"] == "partial"
    assert "dropped 2 rows with no requirement title" in " ".join(
        env["limitations"])


def test_landing_fetch_failure_returns_empty_list_never_raises(caplog):
    def exploding(url, **kw):
        raise ConnectionError("tls reset by waf")

    src = StateForecastSource(fetch_text=exploding)
    with caplog.at_level(logging.WARNING):
        assert src.search(SourceQuery(keywords=["logistics"])) == []
    assert "state_forecast" in caplog.text
    assert "tls reset by waf" in caplog.text


def test_waf_block_page_bytes_return_empty_list(caplog):
    src = _source(data=b"<html>Technical Difficulties</html>")
    with caplog.at_level(logging.WARNING):
        assert src.search(SourceQuery()) == []
    assert "non-xlsx" in caplog.text


def test_landing_page_without_link_returns_empty_list(caplog):
    src = _source(html="<html>Technical Difficulties</html>")
    with caplog.at_level(logging.WARNING):
        assert src.search(SourceQuery()) == []
    assert "no procurement-forecast .xlsx link" in caplog.text


def test_unrecognized_workbook_layout_returns_empty_list(caplog):
    workbook = openpyxl.Workbook()
    workbook.active.append(["Wrong", "Columns"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    src = _source(data=buffer.getvalue())
    with caplog.at_level(logging.WARNING):
        assert src.search(SourceQuery()) == []
    assert "unrecognized forecast workbook layout" in caplog.text


# ------------------------------------------------------------------ #
# healthcheck and registration posture
# ------------------------------------------------------------------ #

def test_healthcheck_failure_shape():
    def dead(url, **kw):
        raise ConnectionError("connect timeout to state.gov")

    ok, detail = StateForecastSource(fetch_text=dead).healthcheck()
    assert ok is False
    assert "unreachable" in detail
    assert "connect timeout" in detail


def test_healthcheck_reports_block_page_honestly():
    src = StateForecastSource(
        fetch_text=lambda url, **kw: "<html>Technical Difficulties</html>")
    ok, detail = src.healthcheck()
    assert ok is False
    assert "no procurement-forecast .xlsx link" in detail


def test_healthcheck_success_names_the_current_workbook():
    src = StateForecastSource(fetch_text=lambda url, **kw: LANDING_HTML)
    ok, detail = src.healthcheck()
    assert ok is True
    assert "FY26-Procurement-Forecast-4.xlsx" in detail


def test_registration_follows_catalog_membership():
    """Deterministic both before and after the lead lands the SourceSpec row:
    registered exactly when cataloged, silent otherwise."""
    from tools.api.source_catalog import SOURCE_BY_ADAPTER
    registered = {s.name for s in REGISTRY.all()}
    assert StateForecastSource.name == "state_forecast"
    assert StateForecastSource.kind is SourceKind.ENRICHMENT
    if "state_forecast" in SOURCE_BY_ADAPTER:
        assert "state_forecast" in registered
        assert REGISTRY.get("state_forecast").kind is SourceKind.ENRICHMENT
    else:
        assert "state_forecast" not in registered


# ------------------------------------------------------------------ #
# opt-in live smoke (the only networked test; smallest possible query)
# ------------------------------------------------------------------ #

@pytest.mark.skipif(not os.environ.get("LILA_LIVE_SMOKE"),
                    reason="live smoke is opt-in: set LILA_LIVE_SMOKE=1")
def test_live_smoke_smallest_query():
    src = StateForecastSource()
    ok, detail = src.healthcheck()
    assert ok, detail
    records = src.search(SourceQuery(naics_codes=["5416"], limit=3))
    assert len(records) <= 3
    for opp in records:
        assert opp.source == "state_forecast"
        assert opp.agency == "Department of State"
        assert opp.source_id.startswith("state-fcst-")
        env = opp.raw_payload.get("_provenance", {})
        assert env.get("source") == "state_forecast"
        assert env.get("retrieval_mode") == "live"
