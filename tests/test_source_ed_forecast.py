"""Department of Education Procurement Forecast adapter: fixture-driven, no network.

Fixture provenance (recorded, not invented):
    tests/fixtures/sources/ed_forecast.json was recorded 2026-08-08 from live
    ed.gov. The workbook was fetched at 2026-08-08T17:19:19Z from
    https://www.ed.gov/media/document/us-department-of-education-procurement-forecast-july-23-2026-114098.xlsx
    (HTTP 200 anonymous, 46,471 bytes, Last-Modified Thu, 30 Jul 2026
    18:40:11 GMT, single sheet 'FY2026', 373 rows after the header). The
    landing-page anchor excerpt was captured in the same session from
    https://www.ed.gov/about/doing-business-ed/contract-opportunities/forecast-of-ed-contract-opportunities
    (HTTP 200 anonymous). The fixture keeps real bytes throughout: the two
    preamble cells (the 'as of July 23, 2026' banner and 'Dept. Funding
    Code: 9100'), the verbatim 14-column header, the first 14 data rows in
    file order (8(a) rows, stated incumbents, a Recompete row, every value
    band shape including the floorless '< $15K'), and the two real trailing
    note rows (ED's disclaimer, first column only).
"""

from __future__ import annotations

import io
import json
import logging
import os
from pathlib import Path

import pytest
from openpyxl import Workbook

import tools.api.ed_forecast as ef
from tools.api.base import SourceKind, SourceQuery
from tools.api.ed_forecast import (
    EdForecastSource,
    _normalize,
    _value_lower_bound,
    discover_workbook_url,
    parse_meta,
)

FIXTURE = Path(__file__).parent / "fixtures" / "sources" / "ed_forecast.json"
WORKBOOK_URL = (
    "https://www.ed.gov/media/document/"
    "us-department-of-education-procurement-forecast-july-23-2026-114098.xlsx"
)


def _fixture() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def _parsed(fx: dict) -> dict:
    """The fixture in the exact shape _parse_workbook returns."""
    return {"sheet": fx["sheet"], "preamble": fx["preamble"], "rows": fx["rows"]}


def _landing_html(fx: dict) -> str:
    return "<html><body>" + fx["landing_html_excerpt"] + "</body></html>"


def _xlsx_bytes(fx: dict, extra_rows=None) -> bytes:
    """A real workbook rebuilt from the RECORDED cells, matching the live
    layout: banner row, funding-code row, header, data rows, one fully blank
    row, then the two trailing note rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = fx["sheet"]
    header = fx["header"]
    for cell in fx["preamble"]:
        ws.append([cell])
    ws.append(header)
    data = [r for r in fx["rows"] if r.get("Contract Name")]
    notes = [r for r in fx["rows"] if not r.get("Contract Name")]
    for row in data + list(extra_rows or []):
        ws.append([row.get(name) for name in header])
    ws.append([None] * len(header))
    for row in notes:
        ws.append([row.get(name) for name in header])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _wired(fx: dict, blob=None, resp_headers=None) -> EdForecastSource:
    landing = _landing_html(fx)
    blob = blob if blob is not None else _xlsx_bytes(fx)
    headers = resp_headers if resp_headers is not None else {
        "Last-Modified": fx["_fixture_provenance"]["workbook_last_modified"]
    }
    return EdForecastSource(
        fetch_text=lambda url, **kw: landing,
        fetch_bytes=lambda url, **kw: (blob, headers),
    )


def test_adapter_identity():
    src = EdForecastSource()
    assert src.name == "ed_forecast"
    assert src.kind is SourceKind.ENRICHMENT
    from tools.toggles import is_enabled
    assert src.enabled is is_enabled("ed_forecast", True)


def test_workbook_url_discovered_from_recorded_anchor():
    """The real href is site-relative; discovery must absolutize it."""
    fx = _fixture()
    assert discover_workbook_url(_landing_html(fx)) == WORKBOOK_URL
    assert discover_workbook_url("<html>no link here</html>") is None


def test_parse_meta_reads_the_recorded_preamble():
    fx = _fixture()
    meta = parse_meta(_parsed(fx))
    assert meta == {
        "as_of_date": "2026-07-23",
        "fiscal_year": "FY2026",
        "dept_funding_code": "9100",
    }
    bare = parse_meta({"sheet": "FY2026", "preamble": [], "rows": []})
    assert bare["as_of_date"] is None
    assert bare["fiscal_year"] == "FY2026"  # sheet-title fallback


def test_search_end_to_end_maps_every_field():
    """Full path off recorded cells: landing discovery, xlsx parse (banner
    scan, blank row skipped, note rows skipped), normalize, provenance."""
    fx = _fixture()
    out = _wired(fx).search(SourceQuery())
    assert len(out) == 14

    first = out[0]
    assert first.source == "ed_forecast"
    assert first.source_id == "ed-fcst-fy26-ap-4186"
    assert first.title == "Architecture and Data Modeling Support"
    assert first.agency == "Department of Education"
    assert first.naics_code == "541990"
    assert first.set_aside is None            # 'TBD' is a placeholder
    assert first.posted_date is None          # forecast stage, never a deal
    assert first.response_deadline is None
    assert first.estimated_value == 1_000_000.0
    assert str(first.api_url) == WORKBOOK_URL
    assert len(first.contacts) == 1
    poc = first.contacts[0]
    assert poc.name == "Young Choi"
    assert poc.email == "YOUNG.CHOI@ED.GOV"
    assert poc.contact_type == "primary"

    payload = first.raw_payload
    assert payload["tracking_no"] == "FY26-AP-4186"
    assert payload["funding_office"] == "FSA"
    assert payload["contracting_office"] == "FSA"
    assert payload["requirement_type"] == "New"
    assert payload["naics_raw"] == 541990
    assert payload["naics_codes_all"] == ["541990"]
    assert payload["type_of_competition"] == "TBD"      # verbatim ride-along
    assert payload["incumbent_contractor_name"] == "TBD"
    assert payload["estimated_value_range"] == ">= $1M and < $5M"
    assert payload["estimated_current_fy_range"] == ">= $1M and < $5M"
    assert payload["target_award_quarter"] == "Q4"
    assert payload["fiscal_year"] == "FY2026"
    assert payload["dept_funding_code"] == "9100"
    assert payload["source_url"] == WORKBOOK_URL
    assert payload["landing_url"] == ef.LANDING_URL

    by_id = {opp.source_id: opp for opp in out}
    eight_a = by_id["ed-fcst-fy26-ap-4213"]
    assert eight_a.set_aside == "8(a)"
    assert eight_a.raw_payload["incumbent_contractor_name"] == "Areeva Solutions, LLC"
    assert by_id["ed-fcst-fy26-ap-4078"].estimated_value == 50_000_000.0
    assert by_id["ed-fcst-fy26-ap-4473"].estimated_value is None  # '< $15K'


def test_provenance_envelope_on_every_payload():
    fx = _fixture()
    out = _wired(fx).search(SourceQuery())
    assert out and all("_provenance" in opp.raw_payload for opp in out)
    env = out[0].raw_payload["_provenance"]
    assert env["schema_version"] == 1
    assert env["source"] == "ed_forecast"
    assert env["status"] == "complete"
    assert env["mode"] == "live_ed_forecast_xlsx"
    assert env["retrieval_mode"] == "live"
    assert env["record_count"] == 14
    assert env["retrieved_at"]
    # the banner's own 'as of' date wins over Last-Modified (2026-07-30)
    assert env["data_as_of"] == "2026-07-23"
    assert "114098.xlsx" in env["public_detail"]
    assert any("stated agency intent" in note for note in env["limitations"])


def test_last_modified_fallback_when_banner_has_no_date():
    fx = _fixture()
    fx["preamble"] = ["FY 2026 forecast", "Dept. Funding Code: 9100"]
    out = _wired(fx).search(SourceQuery())
    env = out[0].raw_payload["_provenance"]
    assert env["data_as_of"] == "2026-07-30"


def test_query_translation_against_recorded_rows():
    """Adapters translate the normalized query; counts are pinned to the
    recorded first-14 rows."""
    fx = _fixture()
    parsed = _parsed(fx)

    def ids(query):
        matched, _, _ = _normalize(parsed, WORKBOOK_URL, query)
        return [opp.source_id for opp in matched]

    assert len(ids(SourceQuery(naics_codes=["5415"]))) == 8   # prefix match
    assert len(ids(SourceQuery(naics_codes=["561492"]))) == 3
    assert len(ids(SourceQuery(agencies=["OFO"]))) == 7       # funding office
    assert len(ids(SourceQuery(agencies=["Education"]))) == 14
    assert ids(SourceQuery(set_asides=["8(a)"])) == [
        "ed-fcst-fy26-ap-4213", "ed-fcst-fy26-ap-3898",
    ]
    # keyword lane covers title AND requirement type AND incumbent name
    assert ids(SourceQuery(keywords=["recompete"])) == [
        "ed-fcst-fy26-ap-3114", "ed-fcst-fy26-ap-3898",
        "ed-fcst-fy26-ap-3026", "ed-fcst-fy26-ap-3649",
    ]
    assert ids(SourceQuery(keywords=["carasoft"])) == ["ed-fcst-fy26-ap-4078"]
    assert ids(SourceQuery(naics_codes=["5415"], set_asides=["8(a)"])) == [
        "ed-fcst-fy26-ap-4213", "ed-fcst-fy26-ap-3898",
    ]


def test_empty_result_is_calm_empty_list():
    fx = _fixture()
    out = _wired(fx).search(SourceQuery(keywords=["zz never a match zz"]))
    assert out == []


def test_note_rows_skipped_malformed_dropped_never_crash():
    """The file's own trailing disclaimer/note rows are furniture, never
    'dropped data'; a genuinely malformed row (data but no contract name)
    drops with a count. (Synthetic inline row: a unit input for the
    tolerance path, distinct from the recorded fixture.)"""
    fx = _fixture()
    parsed = _parsed(fx)
    matched, dropped, notes = _normalize(parsed, WORKBOOK_URL, SourceQuery())
    assert len(matched) == 14
    assert dropped == 0
    assert notes == 2

    broken = {name: None for name in fx["header"]}
    broken["Tracking No."] = "FY26-AP-9999"
    broken["Funding Office"] = "FSA"
    junk_types = dict(fx["rows"][0])
    junk_types["Tracking No."] = "FY26-AP-9998"
    junk_types["Primary NAICS Code (Subject to Change)"] = "TBD"
    junk_types["Estimated Value of Contract $ Range"] = "TBD"
    parsed_plus = {
        "sheet": fx["sheet"],
        "preamble": fx["preamble"],
        "rows": fx["rows"] + [broken, junk_types],
    }
    matched, dropped, notes = _normalize(parsed_plus, WORKBOOK_URL, SourceQuery())
    assert dropped == 1
    assert notes == 2
    survivor = {opp.source_id: opp for opp in matched}["ed-fcst-fy26-ap-9998"]
    assert survivor.naics_code is None
    assert survivor.estimated_value is None


def test_malformed_row_marks_envelope_partial():
    fx = _fixture()
    broken = {name: None for name in fx["header"]}
    broken["Tracking No."] = "FY26-AP-9999"
    broken["Funding Office"] = "FSA"
    out = _wired(fx, blob=_xlsx_bytes(fx, extra_rows=[broken])).search(SourceQuery())
    assert len(out) == 14
    env = out[0].raw_payload["_provenance"]
    assert env["status"] == "partial"
    assert any("dropped 1 rows" in note for note in env["limitations"])


def test_cap_disclosed_in_limitations():
    fx = _fixture()
    out = _wired(fx).search(SourceQuery(limit=3))
    assert len(out) == 3
    assert out[0].source_id == "ed-fcst-fy26-ap-4186"  # arrival order kept
    env = out[0].raw_payload["_provenance"]
    assert env["record_count"] == 3
    assert any("kept first 3 of 14" in note for note in env["limitations"])


def test_id_stability_duplicates_and_hash_fallback():
    """Ids are assigned before filtering: a narrower query never shifts
    which repeated tracking number owns which ordinal. (Synthetic rows:
    the recorded file has zero duplicate tracking numbers.)"""
    fx = _fixture()
    header = fx["header"]

    def row(tracking, name):
        base = {key: None for key in header}
        base["Tracking No."] = tracking
        base["Funding Office"] = "FSA"
        base["Contract Name"] = name
        base["Target Award Quarter"] = "Q1"
        return base

    parsed = {
        "sheet": fx["sheet"],
        "preamble": fx["preamble"],
        "rows": [
            row("FY26-AP-0001", "First Line"),
            row("FY26-AP-0001", "Second Line Distinct"),
            row("not a tracking id", "Hash Fallback Line"),
        ],
    }
    matched, _, _ = _normalize(parsed, WORKBOOK_URL, SourceQuery())
    ids = [opp.source_id for opp in matched]
    assert ids[0] == "ed-fcst-fy26-ap-0001"
    assert ids[1] == "ed-fcst-fy26-ap-0001-2"
    assert ids[2].startswith("ed-fcst-")
    assert len(ids[2]) == len("ed-fcst-") + 16
    narrowed, _, _ = _normalize(
        parsed, WORKBOOK_URL, SourceQuery(keywords=["distinct"])
    )
    assert [opp.source_id for opp in narrowed] == ["ed-fcst-fy26-ap-0001-2"]


def test_value_lower_bound_witnessed_bands():
    """Every band shape in the July 23, 2026 workbook, plus placeholders."""
    assert _value_lower_bound(">= $1M and < $5M") == 1_000_000.0
    assert _value_lower_bound(">= $10M and < $20M") == 10_000_000.0
    assert _value_lower_bound(">= $350K and < $1M") == 350_000.0
    assert _value_lower_bound(">= $15K and < $350K") == 15_000.0
    assert _value_lower_bound(">= $50M") == 50_000_000.0
    assert _value_lower_bound("< $15K") is None   # no stated floor
    assert _value_lower_bound("TBD") is None
    assert _value_lower_bound(None) is None


def test_failure_isolation_returns_empty_and_logs(caplog):
    fx = _fixture()

    def boom(url, **kw):
        raise ConnectionError("dns failure reaching www.ed.gov")

    src = EdForecastSource(fetch_text=boom, fetch_bytes=boom)
    with caplog.at_level(logging.WARNING, logger="tools.api.ed_forecast"):
        assert src.search(SourceQuery()) == []
    assert any("ed_forecast search failed" in rec.getMessage()
               for rec in caplog.records)

    # landing loads but carries no workbook link
    src = EdForecastSource(
        fetch_text=lambda url, **kw: "<html>redesigned page</html>",
        fetch_bytes=boom,
    )
    assert src.search(SourceQuery()) == []

    # workbook url serves an HTML block page instead of xlsx bytes
    src = _wired(fx, blob=b"<html>maintenance</html>")
    assert src.search(SourceQuery()) == []

    # header drift: a real xlsx that lacks the Tracking No. header row
    wb = Workbook()
    wb.active.append(["Some", "Other", "Layout"])
    buffer = io.BytesIO()
    wb.save(buffer)
    src = _wired(fx, blob=buffer.getvalue())
    with caplog.at_level(logging.WARNING, logger="tools.api.ed_forecast"):
        assert src.search(SourceQuery()) == []
    assert any("Tracking No." in rec.getMessage() for rec in caplog.records)


def test_healthcheck_reports_reality():
    fx = _fixture()
    seen = {}

    def fake_get_text(url, **kw):
        seen["url"], seen["kw"] = url, kw
        return _landing_html(fx)

    ok, detail = EdForecastSource(fetch_text=fake_get_text).healthcheck()
    assert ok is True
    assert "114098.xlsx" in detail
    assert seen["url"] == ef.LANDING_URL
    assert seen["kw"] == {"timeout": 8.0, "retries": 1}

    def boom(url, **kw):
        raise ConnectionError("connect timeout")

    ok, detail = EdForecastSource(fetch_text=boom).healthcheck()
    assert ok is False
    assert "unreachable" in detail and "connect timeout" in detail

    ok, detail = EdForecastSource(
        fetch_text=lambda url, **kw: "<html>no links</html>"
    ).healthcheck()
    assert ok is False
    assert ".xlsx" in detail


@pytest.mark.skipif(not os.environ.get("LILA_LIVE_SMOKE"),
                    reason="live smoke is opt-in: set LILA_LIVE_SMOKE=1")
def test_live_smoke_smallest_probe():
    """Smallest possible live touch: one landing-page GET, no workbook pull."""
    ok, detail = EdForecastSource().healthcheck()
    assert ok, detail
